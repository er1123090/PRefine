#!/usr/bin/env python3
"""Export axis-based analysis treating sampled runs as full-population estimates."""

from __future__ import annotations

import argparse
import gc
import json
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
SCRIPT_DIR = ROOT / "scripts"
for path in (ROOT, SCRIPT_DIR):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

import run_local_full_inference as local_runtime  # noqa: E402
import run_vanilla_batch_sample as population_runtime  # noqa: E402
from export_unified_vanilla_ours_performance import (  # noqa: E402
    EXPERIMENT_ROOT,
    RUNS,
    RunSpec,
)


ANALYSIS_N = 4695
POPULATION_META_PATH = (
    EXPERIMENT_ROOT / "population" / "population.jsonl"
)
DEFAULT_OUTPUT = (
    EXPERIMENT_ROOT
    / "native_full_vanilla5_ours3x5_axis_performance.xlsx"
)

STRATUM_ORDER = (
    "single_easy",
    "single_medium",
    "single_hard",
    "single_conflict-medium",
    "single_conflict-hard",
    "multi_easy",
    "multi_medium",
    "multi_hard",
    "multi_conflict-medium",
    "multi_conflict-hard",
)

PERCENT_HEADERS = {
    "Overall Precision",
    "Overall Recall",
    "Overall F1",
    "Preference Precision",
    "Preference Recall",
    "Preference F1",
    "Preference Exact Match",
    "Non-preference Precision",
    "Non-preference Recall",
    "Non-preference F1",
    "API Error Rate",
    "Parsing Failure Rate",
    "Vanilla Overall F1",
    "Ours Overall F1",
    "Overall F1 Delta",
    "Vanilla Preference EM",
    "Ours Preference EM",
    "Preference EM Delta",
    "Vanilla Parsing Rate",
    "Ours Parsing Rate",
    "Parsing Rate Delta",
    "Group Mean Overall F1",
    "Group Mean Preference EM",
    "Weighted Overall F1",
    "Weighted Preference F1",
    "Weighted Preference Exact Match",
    "Weighted Non-preference F1",
    "Weighted API Error Rate",
    "Weighted Parsing Failure Rate",
}

INTEGER_HEADERS = {
    "Observed N",
    "Analysis N",
    "Actual API Errors",
    "Estimated Full API Errors",
    "Actual Parsing Failures",
    "Estimated Full Parsing Failures",
    "Observed Preference EM Count",
    "Estimated Full Preference EM Count",
    "Observed Total Tokens",
    "Estimated Full Total Tokens",
    "Observed Input Tokens",
    "Estimated Full Input Tokens",
    "Observed Output Tokens",
    "Estimated Full Output Tokens",
    "Observed Reasoning Tokens",
    "Estimated Full Reasoning Tokens",
    "Overall F1 Rank",
    "Preference EM Rank",
    "Clean Rank",
    "Rank Within Inference Model",
    "Rank Within Memory",
    "Observed Stratum N",
    "Assumed Full Stratum N",
    "Observed API Errors",
    "Observed Parsing Failures",
    "Observed Aggregate N",
    "Assumed Full Aggregate N",
}


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def dense_rank(
    rows: Sequence[dict[str, Any]],
    metric: str,
    output: str,
    *,
    eligible: callable | None = None,
) -> None:
    selected = [
        row
        for row in rows
        if eligible is None or bool(eligible(row))
    ]
    values = sorted({float(row[metric]) for row in selected}, reverse=True)
    ranks = {value: rank for rank, value in enumerate(values, start=1)}
    for row in rows:
        if row in selected:
            row[output] = ranks[float(row[metric])]
        else:
            row[output] = None


def projected_count(actual: int, observed_n: int) -> int:
    if observed_n == ANALYSIS_N:
        return actual
    return round(actual / observed_n * ANALYSIS_N)


def projected_tokens(actual: int, observed_n: int) -> int:
    if observed_n == ANALYSIS_N:
        return actual
    return round(actual / observed_n * ANALYSIS_N)


def projected_stratum_count(
    actual: int,
    observed_n: int,
    full_stratum_n: int,
) -> int:
    if observed_n == full_stratum_n:
        return actual
    return round(actual / observed_n * full_stratum_n)


def load_run_predictions(
    spec: RunSpec,
    population: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    if spec.scope == "sample":
        return json.loads(
            (spec.directory / "predictions.json").read_text(
                encoding="utf-8"
            )
        )
    records = local_runtime.read_checkpoint(
        spec.directory / "inference.jsonl"
    )
    if len(records) != ANALYSIS_N:
        raise RuntimeError(
            f"{spec.key}: expected {ANALYSIS_N} records, got {len(records)}"
        )
    return local_runtime.materialize_predictions(population, records)


def load_evaluation_report(spec: RunSpec) -> dict[str, Any]:
    path = spec.directory / "evaluation.json"
    if not path.exists():
        raise FileNotFoundError(f"{spec.key}: missing {path}")
    report = json.loads(path.read_text(encoding="utf-8"))
    required = {"overall", "condition_conflict", "token_usage"}
    missing = required - set(report)
    if missing:
        raise RuntimeError(
            f"{spec.key}: evaluation report missing {sorted(missing)}"
        )
    return report


def overall_row(
    spec: RunSpec,
    report: Mapping[str, Any],
) -> dict[str, Any]:
    metrics = report["overall"]
    usage = report.get("token_usage") or {}
    observed_n = int(metrics["n"])
    projected = observed_n != ANALYSIS_N
    api_errors = int(metrics["api_errors"])
    parse_failures = int(metrics["parsing_failures"])
    pref_em_count = int(metrics["pref"]["exact_match_count"])
    return {
        "Run ID": spec.key,
        "Method": spec.method,
        "Memory Source": spec.memory,
        "Inference Model": spec.inference_model,
        "Inference Setting": spec.setting,
        "Provider / Backend": spec.provider,
        "Evidence": (
            "Projected from proportional N=1,000 sample"
            if projected
            else "Measured on all N=4,695 instances"
        ),
        "Observed N": observed_n,
        "Analysis N": ANALYSIS_N,
        "Overall Precision": float(metrics["overall"]["precision"]),
        "Overall Recall": float(metrics["overall"]["recall"]),
        "Overall F1": float(metrics["overall"]["f1"]),
        "Overall F1 Rank": None,
        "Preference Precision": float(metrics["pref"]["precision"]),
        "Preference Recall": float(metrics["pref"]["recall"]),
        "Preference F1": float(metrics["pref"]["f1"]),
        "Preference Exact Match": float(
            metrics["pref"]["exact_match_rate"]
        ),
        "Observed Preference EM Count": pref_em_count,
        "Estimated Full Preference EM Count": projected_count(
            pref_em_count, observed_n
        ),
        "Preference EM Rank": None,
        "Non-preference Precision": float(
            metrics["nonpref"]["precision"]
        ),
        "Non-preference Recall": float(metrics["nonpref"]["recall"]),
        "Non-preference F1": float(metrics["nonpref"]["f1"]),
        "Actual API Errors": api_errors,
        "Estimated Full API Errors": projected_count(
            api_errors, observed_n
        ),
        "API Error Rate": api_errors / observed_n,
        "Actual Parsing Failures": parse_failures,
        "Estimated Full Parsing Failures": projected_count(
            parse_failures, observed_n
        ),
        "Parsing Failure Rate": float(
            metrics["parsing_failure_rate"]
        ),
        "Clean Rank": None,
        "Observed Input Tokens": int(
            usage.get("input_tokens", 0) or 0
        ),
        "Estimated Full Input Tokens": projected_tokens(
            int(usage.get("input_tokens", 0) or 0), observed_n
        ),
        "Observed Output Tokens": int(
            usage.get("output_tokens", 0) or 0
        ),
        "Estimated Full Output Tokens": projected_tokens(
            int(usage.get("output_tokens", 0) or 0), observed_n
        ),
        "Observed Reasoning Tokens": int(
            usage.get("reasoning_tokens", 0) or 0
        ),
        "Estimated Full Reasoning Tokens": projected_tokens(
            int(usage.get("reasoning_tokens", 0) or 0), observed_n
        ),
        "Observed Total Tokens": int(
            usage.get("total_tokens", 0) or 0
        ),
        "Estimated Full Total Tokens": projected_tokens(
            int(usage.get("total_tokens", 0) or 0), observed_n
        ),
        "Tokens per Observed Instance": (
            int(usage.get("total_tokens", 0) or 0) / observed_n
        ),
        "Result Directory": str(spec.directory),
    }


def stratum_for_meta(row: Mapping[str, Any]) -> str:
    case = str(row["pref_type"])
    if bool(row["conflict"]):
        case = f"conflict-{case}"
    return f"{row['turn']}_{case}"


def collect_metrics() -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
]:
    population_meta = read_jsonl(POPULATION_META_PATH)
    if len(population_meta) != ANALYSIS_N:
        raise RuntimeError(
            f"Expected metadata N={ANALYSIS_N}, got {len(population_meta)}"
        )
    full_stratum_counts: dict[str, int] = defaultdict(int)
    for row in population_meta:
        full_stratum_counts[stratum_for_meta(row)] += 1
    if set(full_stratum_counts) != set(STRATUM_ORDER):
        raise RuntimeError(
            f"Unexpected strata: {sorted(full_stratum_counts)}"
        )

    overall_rows: list[dict[str, Any]] = []
    stratum_rows: list[dict[str, Any]] = []
    for index, spec in enumerate(RUNS, start=1):
        print(f"[{index:02d}/{len(RUNS)}] analyzing {spec.key}", flush=True)
        report = load_evaluation_report(spec)
        overall_rows.append(overall_row(spec, report))

        for stratum in STRATUM_ORDER:
            turn, case = stratum.split("_", 1)
            conflict = case.startswith("conflict-")
            difficulty = case.removeprefix("conflict-")
            condition_key = f"{turn}_{difficulty}"
            conflict_key = "conflict" if conflict else "non_conflict"
            try:
                metrics = report["condition_conflict"][condition_key][
                    conflict_key
                ]
            except KeyError as exc:
                raise RuntimeError(
                    f"{spec.key}: missing {condition_key}/{conflict_key}"
                ) from exc
            observed_n = int(metrics["n"])
            full_stratum_n = int(full_stratum_counts[stratum])
            api_errors = int(metrics["api_errors"])
            parse_failures = int(metrics["parsing_failures"])
            pref_em_count = int(metrics["pref"]["exact_match_count"])
            stratum_rows.append(
                {
                    "Run ID": spec.key,
                    "Method": spec.method,
                    "Memory Source": spec.memory,
                    "Inference Model": spec.inference_model,
                    "Inference Setting": spec.setting,
                    "Provider / Backend": spec.provider,
                    "Stratum": stratum,
                    "Turn": turn,
                    "Difficulty": difficulty,
                    "Conflict": "Yes" if conflict else "No",
                    "Evidence": (
                        "Projected stratum rate"
                        if spec.scope == "sample"
                        else "Measured full stratum"
                    ),
                    "Observed Stratum N": observed_n,
                    "Assumed Full Stratum N": full_stratum_n,
                    "Overall Precision": float(
                        metrics["overall"]["precision"]
                    ),
                    "Overall Recall": float(
                        metrics["overall"]["recall"]
                    ),
                    "Overall F1": float(
                        metrics["overall"]["f1"]
                    ),
                    "Preference Precision": float(
                        metrics["pref"]["precision"]
                    ),
                    "Preference Recall": float(
                        metrics["pref"]["recall"]
                    ),
                    "Preference F1": float(metrics["pref"]["f1"]),
                    "Preference Exact Match": float(
                        metrics["pref"]["exact_match_rate"]
                    ),
                    "Observed Preference EM Count": pref_em_count,
                    "Estimated Full Preference EM Count": (
                        projected_stratum_count(
                            pref_em_count,
                            observed_n,
                            full_stratum_n,
                        )
                    ),
                    "Non-preference Precision": float(
                        metrics["nonpref"]["precision"]
                    ),
                    "Non-preference Recall": float(
                        metrics["nonpref"]["recall"]
                    ),
                    "Non-preference F1": float(
                        metrics["nonpref"]["f1"]
                    ),
                    "Observed API Errors": api_errors,
                    "Estimated Full API Errors": projected_stratum_count(
                        api_errors,
                        observed_n,
                        full_stratum_n,
                    ),
                    "API Error Rate": api_errors / observed_n,
                    "Observed Parsing Failures": parse_failures,
                    "Estimated Full Parsing Failures": (
                        projected_stratum_count(
                            parse_failures,
                            observed_n,
                            full_stratum_n,
                        )
                    ),
                    "Parsing Failure Rate": float(
                        metrics["parsing_failure_rate"]
                    ),
                }
            )

    dense_rank(overall_rows, "Overall F1", "Overall F1 Rank")
    dense_rank(
        overall_rows,
        "Preference Exact Match",
        "Preference EM Rank",
    )
    dense_rank(
        overall_rows,
        "Overall F1",
        "Clean Rank",
        eligible=lambda row: (
            row["Actual API Errors"] == 0
            and row["Actual Parsing Failures"] == 0
        ),
    )
    return overall_rows, stratum_rows


def vanilla_vs_memory_rows(
    overall_rows: Sequence[dict[str, Any]],
) -> list[dict[str, Any]]:
    vanilla = {
        row["Inference Model"]: row
        for row in overall_rows
        if row["Method"] == "Vanilla LLM"
    }
    rows: list[dict[str, Any]] = []
    for ours in overall_rows:
        if ours["Method"] != "Ours-memory":
            continue
        base = vanilla[ours["Inference Model"]]
        rows.append(
            {
                "Inference Model": ours["Inference Model"],
                "Memory Source": ours["Memory Source"],
                "Comparison Meaning": (
                    "Combined memory + reasoning-setting effect"
                ),
                "Reasoning Setting Matched": "No",
                "Vanilla Setting": base["Inference Setting"],
                "Ours Setting": ours["Inference Setting"],
                "Vanilla Overall F1": base["Overall F1"],
                "Ours Overall F1": ours["Overall F1"],
                "Overall F1 Delta": (
                    ours["Overall F1"] - base["Overall F1"]
                ),
                "Vanilla Preference EM": base[
                    "Preference Exact Match"
                ],
                "Ours Preference EM": ours[
                    "Preference Exact Match"
                ],
                "Preference EM Delta": (
                    ours["Preference Exact Match"]
                    - base["Preference Exact Match"]
                ),
                "Vanilla Parsing Rate": base[
                    "Parsing Failure Rate"
                ],
                "Ours Parsing Rate": ours[
                    "Parsing Failure Rate"
                ],
                "Parsing Rate Delta": (
                    ours["Parsing Failure Rate"]
                    - base["Parsing Failure Rate"]
                ),
                "Ours Evidence": ours["Evidence"],
            }
        )
    return rows


def memory_source_rows(
    overall_rows: Sequence[dict[str, Any]],
) -> list[dict[str, Any]]:
    rows = [
        {
            "Inference Model": row["Inference Model"],
            "Memory Source": row["Memory Source"],
            "Evidence": row["Evidence"],
            "Overall F1": row["Overall F1"],
            "Preference F1": row["Preference F1"],
            "Preference Exact Match": row["Preference Exact Match"],
            "Non-preference F1": row["Non-preference F1"],
            "Parsing Failure Rate": row["Parsing Failure Rate"],
            "Rank Within Inference Model": None,
        }
        for row in overall_rows
        if row["Method"] == "Ours-memory"
    ]
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[str(row["Inference Model"])].append(row)
    for group in grouped.values():
        dense_rank(
            group,
            "Preference Exact Match",
            "Rank Within Inference Model",
        )
    return rows


def inference_model_rows(
    overall_rows: Sequence[dict[str, Any]],
) -> list[dict[str, Any]]:
    rows = [
        {
            "Memory Source": row["Memory Source"],
            "Inference Model": row["Inference Model"],
            "Inference Setting": row["Inference Setting"],
            "Evidence": row["Evidence"],
            "Overall F1": row["Overall F1"],
            "Preference F1": row["Preference F1"],
            "Preference Exact Match": row["Preference Exact Match"],
            "Non-preference F1": row["Non-preference F1"],
            "Parsing Failure Rate": row["Parsing Failure Rate"],
            "Rank Within Memory": None,
        }
        for row in overall_rows
        if row["Method"] == "Ours-memory"
    ]
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[str(row["Memory Source"])].append(row)
    for group in grouped.values():
        dense_rank(
            group,
            "Preference Exact Match",
            "Rank Within Memory",
        )
    return rows


def task_axis_summary_rows(
    stratum_rows: Sequence[dict[str, Any]],
) -> list[dict[str, Any]]:
    ours = [
        row for row in stratum_rows if row["Method"] == "Ours-memory"
    ]
    run_order = list(dict.fromkeys(str(row["Run ID"]) for row in ours))
    by_run: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in ours:
        by_run[str(row["Run ID"])].append(row)

    output: list[dict[str, Any]] = []
    axis_values = (
        ("Turn", ("single", "multi")),
        ("Difficulty", ("easy", "medium", "hard")),
        ("Conflict", ("No", "Yes")),
    )
    metric_pairs = (
        ("Weighted Overall F1", "Overall F1"),
        ("Weighted Preference F1", "Preference F1"),
        (
            "Weighted Preference Exact Match",
            "Preference Exact Match",
        ),
        ("Weighted Non-preference F1", "Non-preference F1"),
        ("Weighted API Error Rate", "API Error Rate"),
        (
            "Weighted Parsing Failure Rate",
            "Parsing Failure Rate",
        ),
    )
    for run_id in run_order:
        run_rows = by_run[run_id]
        example = run_rows[0]
        for axis, values in axis_values:
            for value in values:
                selected = [
                    row
                    for row in run_rows
                    if (
                        str(row[axis]) == value
                        and (
                            axis != "Conflict"
                            or row["Difficulty"] in {"medium", "hard"}
                        )
                    )
                ]
                if not selected:
                    raise RuntimeError(
                        f"{run_id}: no rows for {axis}={value}"
                    )
                total_weight = sum(
                    int(row["Assumed Full Stratum N"])
                    for row in selected
                )
                summary = {
                    "Run ID": run_id,
                    "Memory Source": example["Memory Source"],
                    "Inference Model": example["Inference Model"],
                    "Inference Setting": example["Inference Setting"],
                    "Axis": axis,
                    "Axis Value": value,
                    "Scope": (
                        "Medium + hard only"
                        if axis == "Conflict"
                        else "All applicable native strata"
                    ),
                    "Evidence": example["Evidence"],
                    "Observed Aggregate N": sum(
                        int(row["Observed Stratum N"])
                        for row in selected
                    ),
                    "Assumed Full Aggregate N": total_weight,
                }
                for output_metric, source_metric in metric_pairs:
                    summary[output_metric] = sum(
                        float(row[source_metric])
                        * int(row["Assumed Full Stratum N"])
                        for row in selected
                    ) / total_weight
                output.append(summary)
    return output


def self_cross_rows(
    overall_rows: Sequence[dict[str, Any]],
) -> list[dict[str, Any]]:
    open_models = {"Qwen3-8B", "GPT-OSS-20B", "Gemma4-12B"}
    details: list[dict[str, Any]] = []
    for row in overall_rows:
        if (
            row["Method"] != "Ours-memory"
            or row["Inference Model"] not in open_models
        ):
            continue
        pairing = (
            "Self"
            if row["Memory Source"] == row["Inference Model"]
            else "Cross"
        )
        details.append(
            {
                "Row Type": "Run",
                "Memory Source": row["Memory Source"],
                "Inference Model": row["Inference Model"],
                "Pairing": pairing,
                "Overall F1": row["Overall F1"],
                "Preference Exact Match": row[
                    "Preference Exact Match"
                ],
                "Parsing Failure Rate": row[
                    "Parsing Failure Rate"
                ],
                "Group Mean Overall F1": None,
                "Group Mean Preference EM": None,
            }
        )
    for pairing in ("Self", "Cross"):
        group = [row for row in details if row["Pairing"] == pairing]
        details.append(
            {
                "Row Type": "Group average",
                "Memory Source": "—",
                "Inference Model": "—",
                "Pairing": pairing,
                "Overall F1": None,
                "Preference Exact Match": None,
                "Parsing Failure Rate": None,
                "Group Mean Overall F1": mean(
                    float(row["Overall F1"]) for row in group
                ),
                "Group Mean Preference EM": mean(
                    float(row["Preference Exact Match"])
                    for row in group
                ),
            }
        )
    return details


def reliability_rows(
    overall_rows: Sequence[dict[str, Any]],
) -> list[dict[str, Any]]:
    headers = (
        "Run ID",
        "Method",
        "Memory Source",
        "Inference Model",
        "Evidence",
        "Observed N",
        "Analysis N",
        "Actual API Errors",
        "Estimated Full API Errors",
        "API Error Rate",
        "Actual Parsing Failures",
        "Estimated Full Parsing Failures",
        "Parsing Failure Rate",
        "Observed Total Tokens",
        "Tokens per Observed Instance",
        "Estimated Full Total Tokens",
    )
    return [{header: row[header] for header in headers} for row in overall_rows]


def write_table_sheet(
    workbook: Workbook,
    *,
    name: str,
    title: str,
    note: str,
    rows: Sequence[Mapping[str, Any]],
    table_name: str,
    color_metrics: Iterable[str] = (),
) -> None:
    if not rows:
        raise ValueError(f"{name}: rows must not be empty")
    headers = tuple(rows[0].keys())
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(headers),
    )
    sheet.cell(1, 1, title)
    sheet.cell(1, 1).font = Font(size=16, bold=True, color="FFFFFF")
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor="17365D")
    sheet.row_dimensions[1].height = 26
    sheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=len(headers),
    )
    sheet.cell(2, 1, note)
    sheet.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.cell(2, 1).fill = PatternFill("solid", fgColor="D9EAF7")
    sheet.row_dimensions[2].height = 42

    header_row = 4
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    for row_index, row in enumerate(rows, start=header_row + 1):
        for column, header in enumerate(headers, start=1):
            value = row.get(header)
            cell = sheet.cell(row_index, column, value)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=header
                in {
                    "Evidence",
                    "Comparison Meaning",
                    "Vanilla Setting",
                    "Ours Setting",
                    "Result Directory",
                },
            )
            if header in PERCENT_HEADERS:
                cell.number_format = "0.00%"
            elif header in INTEGER_HEADERS:
                cell.number_format = "#,##0"
            elif "Tokens per" in header:
                cell.number_format = "#,##0.00"

    last_row = header_row + len(rows)
    last_column = len(headers)
    table = Table(
        displayName=table_name,
        ref=f"A{header_row}:{get_column_letter(last_column)}{last_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = "E5"

    header_columns = {
        header: column
        for column, header in enumerate(headers, start=1)
    }
    for metric in color_metrics:
        if metric not in header_columns:
            continue
        column = get_column_letter(header_columns[metric])
        sheet.conditional_formatting.add(
            f"{column}{header_row + 1}:{column}{last_row}",
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )

    wide_headers = {
        "Run ID": 23,
        "Method": 15,
        "Memory Source": 17,
        "Inference Model": 20,
        "Inference Setting": 25,
        "Provider / Backend": 22,
        "Evidence": 38,
        "Comparison Meaning": 38,
        "Vanilla Setting": 25,
        "Ours Setting": 25,
        "Result Directory": 58,
        "Stratum": 24,
        "Difficulty": 14,
        "Conflict": 12,
    }
    for column, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = (
            wide_headers.get(header, 19)
        )


def write_readme(workbook: Workbook, output_path: Path) -> None:
    sheet = workbook.create_sheet("README")
    sheet.sheet_view.showGridLines = False
    rows = (
        ("Workbook", output_path.name),
        ("Generated", datetime.now(timezone.utc).isoformat()),
        ("Dataset", str(ROOT / "data" / "MPT_v2_0725.json")),
        ("Analysis population", "All runs are analyzed as N=4,695."),
        (
            "Projection assumption",
            "The six GPT-5/Claude ours-memory runs observed on N=1,000 are "
            "treated as proportional full-population estimates. Rates are "
            "kept unchanged; counts and token totals are extrapolated to N=4,695.",
        ),
        (
            "Evidence boundary",
            "Projected values are estimates, not actual 4,695-request executions. "
            "Observed N and Evidence columns retain that distinction.",
        ),
        (
            "Method-effect caveat",
            "Vanilla and ours-memory reasoning settings differ for every inference "
            "model, so Vanilla_vs_Memory measures the combined memory + reasoning "
            "configuration effect, not a pure memory-only causal effect.",
        ),
        (
            "Primary metric",
            "Preference Exact Match directly measures preference-memory correctness; "
            "Overall F1 measures complete API-call quality.",
        ),
        (
            "Task axes",
            "Ours_Axis_Summary provides weighted Turn/Difficulty/Conflict "
            "summaries. Ours_3x5_Task_Axes contains Memory Source × Inference "
            "Model × single/multi × easy/medium/hard × conflict status. "
            "Vanilla_Task_Axes contains the same task axes for the five "
            "memory-free baselines.",
        ),
        (
            "Failure handling",
            "API and parsing failures remain in denominators and lower performance.",
        ),
        (
            "No Comparable_1000",
            "This workbook intentionally contains no common-1,000 comparison sheet.",
        ),
    )
    for row_index, (key, value) in enumerate(rows, start=1):
        sheet.cell(row_index, 1, key).font = Font(bold=True)
        sheet.cell(row_index, 1).fill = PatternFill(
            "solid", fgColor="D9EAF7"
        )
        sheet.cell(row_index, 2, value).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        sheet.row_dimensions[row_index].height = 34
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 120


def build_workbook(output_path: Path) -> dict[str, Any]:
    overall_rows, stratum_rows = collect_metrics()
    vanilla_rows = vanilla_vs_memory_rows(overall_rows)
    memory_rows = memory_source_rows(overall_rows)
    inference_rows = inference_model_rows(overall_rows)
    task_summary_rows = task_axis_summary_rows(stratum_rows)
    pairing_rows = self_cross_rows(overall_rows)
    reliability = reliability_rows(overall_rows)
    ours_stratum_rows = [
        row for row in stratum_rows if row["Method"] == "Ours-memory"
    ]
    vanilla_stratum_rows = [
        row for row in stratum_rows if row["Method"] == "Vanilla LLM"
    ]

    workbook = Workbook()
    workbook.remove(workbook.active)
    write_table_sheet(
        workbook,
        name="Overall_Native_Pop",
        title="Vanilla 5 + Ours-memory 3×5 — Native Population Analysis",
        note=(
            "All 20 runs are analyzed as N=4,695. Six GPT-5/Claude memory runs "
            "use rate-preserving projections from observed N=1,000; see Evidence."
        ),
        rows=overall_rows,
        table_name="OverallNativePopulation",
        color_metrics=("Overall F1", "Preference Exact Match"),
    )
    write_table_sheet(
        workbook,
        name="Vanilla_vs_Memory",
        title="Axis 1 — Vanilla vs Ours-memory",
        note=(
            "Each memory run is compared with the same inference-model vanilla "
            "baseline. Reasoning settings are not matched, so deltas represent "
            "the combined memory + reasoning-setting effect."
        ),
        rows=vanilla_rows,
        table_name="VanillaMemoryEffect",
        color_metrics=("Overall F1 Delta", "Preference EM Delta"),
    )
    write_table_sheet(
        workbook,
        name="Memory_Source",
        title="Axis 2 — Memory Source Effect",
        note=(
            "For each fixed inference model, Qwen/Gemma/GPT-OSS memory sources "
            "are ranked by Preference Exact Match."
        ),
        rows=memory_rows,
        table_name="MemorySourceEffect",
        color_metrics=("Overall F1", "Preference Exact Match"),
    )
    write_table_sheet(
        workbook,
        name="Inference_Model",
        title="Axis 3 — Inference Model Effect",
        note=(
            "For each fixed memory source, five inference models are ranked by "
            "Preference Exact Match."
        ),
        rows=inference_rows,
        table_name="InferenceModelEffect",
        color_metrics=("Overall F1", "Preference Exact Match"),
    )
    write_table_sheet(
        workbook,
        name="Ours_Axis_Summary",
        title="Ours-memory 3×5 — Native Task-axis Summary",
        note=(
            "Instance-weighted summaries for Turn, Difficulty, and Conflict. "
            "Conflict compares Yes vs No only within medium+hard so easy cases "
            "do not confound the comparison. F1 columns are weighted stratum "
            "means; the detailed exact strata remain in Ours_3x5_Task_Axes."
        ),
        rows=task_summary_rows,
        table_name="OursAxisSummary",
        color_metrics=(
            "Weighted Overall F1",
            "Weighted Preference Exact Match",
        ),
    )
    write_table_sheet(
        workbook,
        name="Self_vs_Cross",
        title="Axis 4 — Self-memory vs Cross-memory",
        note=(
            "The 3×3 open-model block is labeled Self or Cross. Group averages "
            "are unweighted because every analyzed run assumes N=4,695."
        ),
        rows=pairing_rows,
        table_name="SelfCrossEffect",
        color_metrics=(
            "Overall F1",
            "Preference Exact Match",
            "Group Mean Overall F1",
            "Group Mean Preference EM",
        ),
    )
    write_table_sheet(
        workbook,
        name="Ours_3x5_Task_Axes",
        title="Ours-memory 3×5 — Turn, Difficulty, and Conflict",
        note=(
            "Primary analysis table: Memory Source × Inference Model × "
            "single/multi × easy/medium/hard × conflict status. GPT-5/Claude "
            "sample rates are projected to each native full-population stratum."
        ),
        rows=ours_stratum_rows,
        table_name="OursTaskAxesPerformance",
        color_metrics=("Overall F1", "Preference Exact Match"),
    )
    write_table_sheet(
        workbook,
        name="Vanilla_Task_Axes",
        title="Vanilla 5 — Turn, Difficulty, and Conflict",
        note=(
            "Memory-free native baselines on the same ten task strata. "
            "Memory Source is None."
        ),
        rows=vanilla_stratum_rows,
        table_name="VanillaTaskAxesPerformance",
        color_metrics=("Overall F1", "Preference Exact Match"),
    )
    write_table_sheet(
        workbook,
        name="Reliability_Efficiency",
        title="Axes 8–10 — Reliability and Token Efficiency",
        note=(
            "Actual observed failures are preserved. Estimated full counts and "
            "tokens extrapolate sampled runs linearly to N=4,695."
        ),
        rows=reliability,
        table_name="ReliabilityEfficiency",
        color_metrics=(),
    )
    write_readme(workbook, output_path)
    workbook.active = 0
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    verified = load_workbook(output_path, read_only=False, data_only=False)
    expected = (
        "Overall_Native_Pop",
        "Vanilla_vs_Memory",
        "Memory_Source",
        "Inference_Model",
        "Ours_Axis_Summary",
        "Self_vs_Cross",
        "Ours_3x5_Task_Axes",
        "Vanilla_Task_Axes",
        "Reliability_Efficiency",
        "README",
    )
    if tuple(verified.sheetnames) != expected:
        raise RuntimeError(
            f"Unexpected workbook sheets: {verified.sheetnames}"
        )
    if verified["Overall_Native_Pop"].max_row != 24:
        raise RuntimeError("Overall sheet must contain 20 runs")
    if verified["Vanilla_vs_Memory"].max_row != 19:
        raise RuntimeError("Vanilla comparison must contain 15 rows")
    if verified["Ours_Axis_Summary"].max_row != 109:
        raise RuntimeError("Ours axis summary must contain 105 rows")
    if verified["Ours_3x5_Task_Axes"].max_row != 154:
        raise RuntimeError("Ours task axes must contain 150 rows")
    if verified["Vanilla_Task_Axes"].max_row != 54:
        raise RuntimeError("Vanilla task axes must contain 50 rows")
    verified.close()

    best_f1 = max(overall_rows, key=lambda row: row["Overall F1"])
    best_pref = max(
        overall_rows, key=lambda row: row["Preference Exact Match"]
    )
    clean = [
        row
        for row in overall_rows
        if row["Actual API Errors"] == 0
        and row["Actual Parsing Failures"] == 0
    ]
    best_clean = max(clean, key=lambda row: row["Overall F1"])
    return {
        "output": str(output_path),
        "run_count": len(overall_rows),
        "projected_run_count": sum(
            row["Observed N"] != ANALYSIS_N for row in overall_rows
        ),
        "ours_task_strata_rows": len(ours_stratum_rows),
        "vanilla_task_strata_rows": len(vanilla_stratum_rows),
        "best_overall_f1": {
            "run_id": best_f1["Run ID"],
            "value": best_f1["Overall F1"],
        },
        "best_preference_em": {
            "run_id": best_pref["Run ID"],
            "value": best_pref["Preference Exact Match"],
        },
        "best_clean_overall_f1": {
            "run_id": best_clean["Run ID"],
            "value": best_clean["Overall F1"],
        },
        "sheets": list(expected),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    summary = build_workbook(args.output.resolve())
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
