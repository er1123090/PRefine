#!/usr/bin/env python3
"""Export Experiment8 full-6508 performance comparisons to a styled workbook."""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = (
    ROOT
    / "outputs"
    / "reports"
    / "qwen3_8b_memory_performance_comparison_6508.xlsx"
)

MODEL_SPECS = (
    {
        "name": "Qwen3-8B Memory-based",
        "method": "Ours memory / high reasoning",
        "path": ROOT
        / "outputs"
        / "ours_memory"
        / "full_6508"
        / "qwen3_8b_memory__to__qwen3_8b_high_reasoning_optimized"
        / "evaluation.json",
        "highlight": True,
    },
    {
        "name": "Qwen3-8B Vanilla",
        "method": "Vanilla / thinking off",
        "path": ROOT
        / "outputs"
        / "vanilla_llm"
        / "full_6508"
        / "qwen_qwen3-8b_thinking_off_optimized"
        / "evaluation.json",
        "highlight": False,
    },
    {
        "name": "Claude Haiku 4.5 Vanilla",
        "method": "Vanilla / non-reasoning",
        "path": ROOT
        / "outputs"
        / "vanilla_llm"
        / "full_6508"
        / "anthropic_claude-haiku-4-5_non_reasoning"
        / "evaluation.json",
        "highlight": False,
    },
    {
        "name": "GPT-5 Minimal Vanilla",
        "method": "Vanilla / minimal reasoning",
        "path": ROOT
        / "outputs"
        / "vanilla_llm"
        / "full_6508"
        / "openai_gpt-5_minimal"
        / "evaluation.json",
        "highlight": False,
    },
)

CONDITION_ORDER = (
    "single_easy",
    "single_medium",
    "single_hard",
    "multi_easy",
    "multi_medium",
    "multi_hard",
)

NAVY = "1F4E78"
BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
GREEN = "70AD47"
LIGHT_GREEN = "E2F0D9"
ORANGE = "ED7D31"
LIGHT_ORANGE = "FCE4D6"
GRAY = "A5A5A5"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"
RED = "C00000"
THIN_GRAY = Side(style="thin", color="D9E1F2")
SECTION_FILL = PatternFill("solid", fgColor=LIGHT_BLUE)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HIGHLIGHT_FILL = PatternFill("solid", fgColor=LIGHT_GREEN)


def read_models() -> list[dict[str, Any]]:
    models: list[dict[str, Any]] = []
    for spec in MODEL_SPECS:
        path = Path(spec["path"])
        if not path.is_file():
            raise FileNotFoundError(f"Missing evaluation file: {path}")
        report = json.loads(path.read_text(encoding="utf-8"))
        if report["overall"]["n"] != 6508:
            raise ValueError(
                f"Expected 6,508 rows in {path}, got {report['overall']['n']}"
            )
        models.append({**spec, "report": report})
    return models


def title(
    ws: Any,
    text: str,
    subtitle: str,
    *,
    last_column: int,
) -> None:
    end = get_column_letter(last_column)
    ws.merge_cells(f"A1:{end}1")
    ws["A1"] = text
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(f"A2:{end}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, italic=True, color="44546A")
    ws["A2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 22


def section(ws: Any, row: int, text: str, *, last_column: int) -> None:
    end = get_column_letter(last_column)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(size=12, bold=True, color=NAVY)
    cell.fill = SECTION_FILL
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 21


def header(ws: Any, row: int, values: list[str]) -> None:
    for column, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=column, value=value)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = Border(bottom=THIN_GRAY)
    ws.row_dimensions[row].height = 34


def add_table(
    ws: Any,
    *,
    name: str,
    start_row: int,
    end_row: int,
    end_column: int,
) -> None:
    ref = f"A{start_row}:{get_column_letter(end_column)}{end_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def percent(cell: Any) -> None:
    cell.number_format = "0.00%"


def integer(cell: Any) -> None:
    cell.number_format = "#,##0"


def decimal(cell: Any) -> None:
    cell.number_format = "#,##0.0"


def autofit(ws: Any, *, minimum: int = 10, maximum: int = 34) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        width = minimum
        for cell in column_cells:
            if cell.value is None:
                continue
            width = max(width, min(maximum, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def finish_sheet(ws: Any, *, freeze: str, landscape: bool = True) -> None:
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5


def model_metrics(model: dict[str, Any]) -> dict[str, Any]:
    report = model["report"]
    overall = report["overall"]
    usage = report.get("token_usage", {})
    total_tokens = int(usage.get("total_tokens", 0) or 0)
    return {
        "model": model["name"],
        "method": model["method"],
        "n": overall["n"],
        "precision": overall["overall"]["precision"],
        "recall": overall["overall"]["recall"],
        "f1": overall["overall"]["f1"],
        "pref_f1": overall["pref"]["f1"],
        "exact_match": overall["pref"]["exact_match_rate"],
        "parse_failures": overall["parsing_failures"],
        "parse_rate": overall["parsing_failure_rate"],
        "api_errors": overall["api_errors"],
        "total_tokens": total_tokens,
        "tokens_per_instance": total_tokens / overall["n"],
    }


def overview_sheet(wb: Workbook, models: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("Overview")
    title(
        ws,
        "Qwen3-8B Memory-based Performance Comparison",
        "Experiment8 · identical 6,508-instance population · current evaluation logic",
        last_column=13,
    )
    metrics = [model_metrics(model) for model in models]
    memory = metrics[0]
    vanilla = metrics[1]

    section(ws, 4, "Key Qwen3-8B Comparison", last_column=4)
    header(ws, 5, ["Metric", "Memory-based", "Vanilla", "Difference"])
    comparisons = (
        ("Overall F1", memory["f1"], vanilla["f1"], memory["f1"] - vanilla["f1"]),
        (
            "Exact Match",
            memory["exact_match"],
            vanilla["exact_match"],
            memory["exact_match"] - vanilla["exact_match"],
        ),
        (
            "Parsing Failures",
            memory["parse_failures"],
            vanilla["parse_failures"],
            memory["parse_failures"] - vanilla["parse_failures"],
        ),
        (
            "Total Tokens",
            memory["total_tokens"],
            vanilla["total_tokens"],
            memory["total_tokens"] - vanilla["total_tokens"],
        ),
    )
    for row_number, values in enumerate(comparisons, start=6):
        for column, value in enumerate(values, start=1):
            ws.cell(row=row_number, column=column, value=value)
        if row_number in (6, 7):
            for column in (2, 3, 4):
                percent(ws.cell(row=row_number, column=column))
        else:
            for column in (2, 3, 4):
                integer(ws.cell(row=row_number, column=column))
        ws.cell(row=row_number, column=2).fill = HIGHLIGHT_FILL
    ws["A11"] = (
        "Result: memory-based improves overall F1 by "
        f"{(memory['f1'] - vanilla['f1']) * 100:.2f} percentage points "
        "with zero API errors."
    )
    ws.merge_cells("A11:M11")
    ws["A11"].font = Font(bold=True, color=NAVY)
    ws["A11"].fill = PatternFill("solid", fgColor=LIGHT_GREEN)

    section(ws, 13, "Full 6,508-instance Baseline Comparison", last_column=13)
    columns = [
        "Model",
        "Method",
        "N",
        "Precision",
        "Recall",
        "Overall F1",
        "Preference F1",
        "Exact Match",
        "Parsing Failures",
        "Parse Failure Rate",
        "API Errors",
        "Total Tokens",
        "Tokens / Instance",
    ]
    header(ws, 14, columns)
    for row_number, values in enumerate(metrics, start=15):
        row = [
            values["model"],
            values["method"],
            values["n"],
            values["precision"],
            values["recall"],
            values["f1"],
            values["pref_f1"],
            values["exact_match"],
            values["parse_failures"],
            values["parse_rate"],
            values["api_errors"],
            values["total_tokens"],
            values["tokens_per_instance"],
        ]
        for column, value in enumerate(row, start=1):
            ws.cell(row=row_number, column=column, value=value)
        for column in (4, 5, 6, 7, 8, 10):
            percent(ws.cell(row=row_number, column=column))
        for column in (3, 9, 11, 12):
            integer(ws.cell(row=row_number, column=column))
        decimal(ws.cell(row=row_number, column=13))
        if row_number == 15:
            for column in range(1, len(columns) + 1):
                ws.cell(row=row_number, column=column).fill = HIGHLIGHT_FILL

    add_table(
        ws,
        name="OverviewComparison",
        start_row=14,
        end_row=14 + len(metrics),
        end_column=len(columns),
    )
    ws.conditional_formatting.add(
        f"F15:F{14 + len(metrics)}",
        ColorScaleRule(
            start_type="min",
            start_color="F8696B",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFEB84",
            end_type="max",
            end_color="63BE7B",
        ),
    )
    ws.conditional_formatting.add(
        f"H15:H{14 + len(metrics)}",
        ColorScaleRule(
            start_type="min",
            start_color="F8696B",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFEB84",
            end_type="max",
            end_color="63BE7B",
        ),
    )

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Overall F1 and Exact Match"
    chart.y_axis.title = "Score"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 1
    chart.x_axis.title = "Model"
    chart.height = 8
    chart.width = 15
    chart.add_data(
        Reference(ws, min_col=6, max_col=8, min_row=14, max_row=18),
        titles_from_data=True,
        from_rows=False,
    )
    chart.set_categories(Reference(ws, min_col=1, min_row=15, max_row=18))
    chart.series.pop(1)
    ws.add_chart(chart, "A21")

    ws["A39"] = (
        "Token counts across different model providers may use different tokenizers. "
        "The Qwen3-8B memory-vs-vanilla token comparison is the strongest like-for-like comparison."
    )
    ws.merge_cells("A39:M39")
    ws["A39"].font = Font(italic=True, color="666666")
    autofit(ws)
    ws.column_dimensions["A"].width = 29
    ws.column_dimensions["B"].width = 29
    finish_sheet(ws, freeze="A14")


def condition_sheet(wb: Workbook, models: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("Condition Comparison")
    title(
        ws,
        "Condition-level Performance",
        "F1 and preference exact-match results by single/multi and difficulty",
        last_column=12,
    )
    columns = [
        "Condition",
        "N",
        "Memory F1",
        "Qwen Vanilla F1",
        "F1 Difference",
        "Claude F1",
        "GPT-5 F1",
        "Memory Exact Match",
        "Qwen Vanilla Exact Match",
        "Exact Match Difference",
        "Claude Exact Match",
        "GPT-5 Exact Match",
    ]
    header(ws, 4, columns)
    reports = [model["report"] for model in models]
    for row_number, condition in enumerate(CONDITION_ORDER, start=5):
        condition_rows = [report["conditions"][condition] for report in reports]
        row = [
            condition,
            condition_rows[0]["n"],
            condition_rows[0]["overall"]["f1"],
            condition_rows[1]["overall"]["f1"],
            condition_rows[0]["overall"]["f1"]
            - condition_rows[1]["overall"]["f1"],
            condition_rows[2]["overall"]["f1"],
            condition_rows[3]["overall"]["f1"],
            condition_rows[0]["pref"]["exact_match_rate"],
            condition_rows[1]["pref"]["exact_match_rate"],
            condition_rows[0]["pref"]["exact_match_rate"]
            - condition_rows[1]["pref"]["exact_match_rate"],
            condition_rows[2]["pref"]["exact_match_rate"],
            condition_rows[3]["pref"]["exact_match_rate"],
        ]
        for column, value in enumerate(row, start=1):
            ws.cell(row=row_number, column=column, value=value)
        integer(ws.cell(row=row_number, column=2))
        for column in range(3, 13):
            percent(ws.cell(row=row_number, column=column))
        for column in (3, 5, 8, 10):
            ws.cell(row=row_number, column=column).fill = HIGHLIGHT_FILL
        if condition == "single_hard":
            ws.cell(row=row_number, column=1).fill = PatternFill(
                "solid", fgColor=LIGHT_ORANGE
            )
            ws.cell(row=row_number, column=3).font = Font(
                bold=True, color=RED
            )
    add_table(
        ws,
        name="ConditionComparison",
        start_row=4,
        end_row=4 + len(CONDITION_ORDER),
        end_column=len(columns),
    )
    for target in ("C5:C10", "H5:H10"):
        ws.conditional_formatting.add(
            target,
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )

    f1_chart = BarChart()
    f1_chart.type = "col"
    f1_chart.style = 10
    f1_chart.title = "Memory-based vs Qwen Vanilla F1"
    f1_chart.y_axis.title = "F1"
    f1_chart.y_axis.scaling.min = 0
    f1_chart.y_axis.scaling.max = 1
    f1_chart.height = 8
    f1_chart.width = 16
    f1_chart.add_data(
        Reference(ws, min_col=3, max_col=4, min_row=4, max_row=10),
        titles_from_data=True,
    )
    f1_chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=10))
    ws.add_chart(f1_chart, "A13")

    em_chart = BarChart()
    em_chart.type = "col"
    em_chart.style = 11
    em_chart.title = "Memory-based vs Qwen Vanilla Exact Match"
    em_chart.y_axis.title = "Exact Match"
    em_chart.y_axis.scaling.min = 0
    em_chart.y_axis.scaling.max = 1
    em_chart.height = 8
    em_chart.width = 16
    em_chart.add_data(
        Reference(ws, min_col=8, max_col=9, min_row=4, max_row=10),
        titles_from_data=True,
    )
    em_chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=10))
    ws.add_chart(em_chart, "A30")

    ws["A47"] = (
        "Primary bottleneck: single_hard has 17.00% F1 and gains only "
        "1.31 percentage points over Qwen3-8B vanilla."
    )
    ws.merge_cells("A47:L47")
    ws["A47"].font = Font(bold=True, color=RED)
    ws["A47"].fill = PatternFill("solid", fgColor=LIGHT_ORANGE)
    autofit(ws)
    ws.column_dimensions["A"].width = 20
    finish_sheet(ws, freeze="C5")


def detailed_sheet(wb: Workbook, models: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("Detailed Metrics")
    title(
        ws,
        "Detailed Evaluation Metrics",
        "Flattened overall and condition-level metrics for downstream analysis",
        last_column=18,
    )
    columns = [
        "Model",
        "Scope",
        "N",
        "API Errors",
        "Parsing Failures",
        "Parse Failure Rate",
        "Overall Precision",
        "Overall Recall",
        "Overall F1",
        "Preference Precision",
        "Preference Recall",
        "Preference F1",
        "Exact Match Rate",
        "Exact Match Count",
        "Exact Match Total",
        "Nonpref Precision",
        "Nonpref Recall",
        "Nonpref F1",
    ]
    header(ws, 4, columns)
    row_number = 5
    for model in models:
        report = model["report"]
        scopes = [("overall", report["overall"])]
        scopes.extend(
            (condition, report["conditions"][condition])
            for condition in CONDITION_ORDER
        )
        for scope, values in scopes:
            row = [
                model["name"],
                scope,
                values["n"],
                values["api_errors"],
                values["parsing_failures"],
                values["parsing_failure_rate"],
                values["overall"]["precision"],
                values["overall"]["recall"],
                values["overall"]["f1"],
                values["pref"]["precision"],
                values["pref"]["recall"],
                values["pref"]["f1"],
                values["pref"]["exact_match_rate"],
                values["pref"]["exact_match_count"],
                values["pref"]["total"],
                values["nonpref"]["precision"],
                values["nonpref"]["recall"],
                values["nonpref"]["f1"],
            ]
            for column, value in enumerate(row, start=1):
                ws.cell(row=row_number, column=column, value=value)
            for column in (3, 4, 5, 14, 15):
                integer(ws.cell(row=row_number, column=column))
            for column in (6, 7, 8, 9, 10, 11, 12, 13, 16, 17, 18):
                percent(ws.cell(row=row_number, column=column))
            if model["highlight"]:
                for column in range(1, len(columns) + 1):
                    ws.cell(row=row_number, column=column).fill = HIGHLIGHT_FILL
            row_number += 1
    add_table(
        ws,
        name="DetailedMetrics",
        start_row=4,
        end_row=row_number - 1,
        end_column=len(columns),
    )
    ws.conditional_formatting.add(
        f"I5:I{row_number - 1}",
        ColorScaleRule(
            start_type="min",
            start_color="F8696B",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFEB84",
            end_type="max",
            end_color="63BE7B",
        ),
    )
    autofit(ws)
    ws.column_dimensions["A"].width = 29
    finish_sheet(ws, freeze="C5")


def token_sheet(wb: Workbook, models: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("Token Usage")
    title(
        ws,
        "Recorded Token Usage",
        "Provider-reported token usage; cross-provider tokenizer differences apply",
        last_column=12,
    )
    columns = [
        "Model",
        "N",
        "Total Tokens",
        "Input Tokens",
        "Output Tokens",
        "Reasoning Tokens",
        "Cached Input Tokens",
        "Cache Creation Input Tokens",
        "Total / Instance",
        "Input / Instance",
        "Output / Instance",
        "vs Qwen Vanilla Total",
    ]
    header(ws, 4, columns)
    qwen_vanilla_total = models[1]["report"]["token_usage"]["total_tokens"]
    for row_number, model in enumerate(models, start=5):
        report = model["report"]
        usage = report["token_usage"]
        n = report["overall"]["n"]
        total = int(usage.get("total_tokens", 0) or 0)
        input_tokens = int(usage.get("input_tokens", 0) or 0)
        output_tokens = int(usage.get("output_tokens", 0) or 0)
        row = [
            model["name"],
            n,
            total,
            input_tokens,
            output_tokens,
            int(usage.get("reasoning_tokens", 0) or 0),
            int(usage.get("cached_input_tokens", 0) or 0),
            int(usage.get("cache_creation_input_tokens", 0) or 0),
            total / n,
            input_tokens / n,
            output_tokens / n,
            total / qwen_vanilla_total - 1,
        ]
        for column, value in enumerate(row, start=1):
            ws.cell(row=row_number, column=column, value=value)
        for column in range(2, 9):
            integer(ws.cell(row=row_number, column=column))
        for column in (9, 10, 11):
            decimal(ws.cell(row=row_number, column=column))
        percent(ws.cell(row=row_number, column=12))
        if model["highlight"]:
            for column in range(1, len(columns) + 1):
                ws.cell(row=row_number, column=column).fill = HIGHLIGHT_FILL
    add_table(
        ws,
        name="TokenUsage",
        start_row=4,
        end_row=4 + len(models),
        end_column=len(columns),
    )

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Total Recorded Tokens"
    chart.y_axis.title = "Tokens"
    chart.height = 8
    chart.width = 16
    chart.add_data(
        Reference(ws, min_col=3, min_row=4, max_row=8),
        titles_from_data=True,
    )
    chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=8))
    ws.add_chart(chart, "A11")
    ws["A28"] = (
        "Qwen3-8B memory-based records 15.7% fewer total tokens than "
        "Qwen3-8B vanilla on the same 6,508 instances."
    )
    ws.merge_cells("A28:L28")
    ws["A28"].font = Font(bold=True, color=NAVY)
    ws["A28"].fill = HIGHLIGHT_FILL
    autofit(ws)
    ws.column_dimensions["A"].width = 29
    finish_sheet(ws, freeze="B5")


def sources_sheet(wb: Workbook, models: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("Sources & Notes")
    title(
        ws,
        "Sources and Measurement Notes",
        "Reproducibility metadata for this workbook",
        last_column=4,
    )
    section(ws, 4, "Source Evaluation Files", last_column=4)
    header(ws, 5, ["Model", "Method", "Evaluation File", "N"])
    for row_number, model in enumerate(models, start=6):
        path = Path(model["path"]).resolve()
        values = [
            model["name"],
            model["method"],
            str(path),
            model["report"]["overall"]["n"],
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row=row_number, column=column, value=value)
        integer(ws.cell(row=row_number, column=4))
    add_table(
        ws,
        name="SourceFiles",
        start_row=5,
        end_row=5 + len(models),
        end_column=4,
    )
    section(ws, 12, "Notes", last_column=4)
    notes = (
        "Generated from Experiment8 evaluation.json files.",
        "All compared evaluation files contain exactly 6,508 instances.",
        "Overall F1 is the current Experiment8 evaluator's aggregate F1.",
        "Exact Match is preference exact-match rate.",
        "Parsing failures are counted separately from API errors.",
        "Token counts are provider-reported; cross-provider tokenizers differ.",
        "The like-for-like token comparison is Qwen3-8B memory-based versus Qwen3-8B vanilla.",
        "The workbook does not modify predictions or evaluation results.",
    )
    for row_number, note in enumerate(notes, start=13):
        ws.cell(row=row_number, column=1, value=f"• {note}")
        ws.merge_cells(
            start_row=row_number,
            start_column=1,
            end_row=row_number,
            end_column=4,
        )
    ws["A23"] = "Generated at"
    ws["B23"] = datetime.now().astimezone().isoformat(timespec="seconds")
    ws["A24"] = "Generator"
    ws["B24"] = str(Path(__file__).resolve())
    ws["A23"].font = ws["A24"].font = Font(bold=True, color=NAVY)
    autofit(ws, maximum=90)
    ws.column_dimensions["A"].width = 31
    ws.column_dimensions["B"].width = 31
    ws.column_dimensions["C"].width = 90
    finish_sheet(ws, freeze="A6", landscape=False)


def build_workbook(output: Path) -> None:
    models = read_models()
    workbook = Workbook()
    workbook.remove(workbook.active)
    overview_sheet(workbook, models)
    condition_sheet(workbook, models)
    detailed_sheet(workbook, models)
    token_sheet(workbook, models)
    sources_sheet(workbook, models)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)

    reopened = load_workbook(output, read_only=False, data_only=False)
    expected_sheets = {
        "Overview",
        "Condition Comparison",
        "Detailed Metrics",
        "Token Usage",
        "Sources & Notes",
    }
    if set(reopened.sheetnames) != expected_sheets:
        raise RuntimeError(f"Workbook sheet validation failed: {reopened.sheetnames}")
    expected_f1 = models[0]["report"]["overall"]["overall"]["f1"]
    workbook_f1 = reopened["Overview"]["F15"].value
    if workbook_f1 is None or abs(workbook_f1 - expected_f1) > 1e-15:
        raise RuntimeError("Workbook metric validation failed")
    reopened.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    build_workbook(args.output.resolve())
    print(args.output.resolve())
