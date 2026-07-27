#!/usr/bin/env python3
"""Export the MPT_v2_0725 vanilla-LLM evaluation to a detailed Excel workbook.

The workbook is generated from the current evaluation.json files and the
matching inference/population JSONL files.  It intentionally contains both
aggregate tables and auditable per-instance rows.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from itertools import combinations
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_EXPERIMENT_ROOT = (
    ROOT / "outputs" / "mpt_v2_0725_hint_no_easy_conflict"
)
DEFAULT_OUTPUT = DEFAULT_EXPERIMENT_ROOT / "vanilla_performance_detailed.xlsx"

CONDITION_ORDER = (
    "single_easy",
    "single_medium",
    "single_hard",
    "multi_easy",
    "multi_medium",
    "multi_hard",
)
TURN_ORDER = ("single", "multi")
DIFFICULTY_ORDER = ("easy", "medium", "hard")
CONFLICT_ORDER = ("non_conflict", "conflict")

MODEL_SPECS = (
    {
        "key": "gpt5",
        "name": "GPT-5",
        "directory": "openai_gpt-5_minimal",
        "provider": "OpenAI",
        "request_model": "gpt-5",
        "reasoning": "minimal",
        "inference_mode": "OpenAI Batch API",
    },
    {
        "key": "claude45",
        "name": "Claude Haiku 4.5",
        "directory": "anthropic_claude-haiku-4-5_non_reasoning",
        "provider": "Anthropic",
        "request_model": "claude-haiku-4-5",
        "reasoning": "thinking disabled",
        "inference_mode": "Anthropic Message Batches API",
    },
    {
        "key": "qwen3_8b",
        "name": "Qwen3-8B",
        "directory": "qwen_qwen3-8b_thinking_off_optimized",
        "provider": "OpenRouter",
        "request_model": "qwen/qwen3-8b",
        "reasoning": "thinking off / reasoning none",
        "inference_mode": "OpenRouter paid API",
    },
    {
        "key": "gpt_oss20b",
        "name": "GPT-OSS-20B",
        "directory": "openai_gpt-oss-20b_thinking_off_reasoning_low_optimized",
        "provider": "OpenRouter",
        "request_model": "openai/gpt-oss-20b",
        "reasoning": "thinking off / reasoning low",
        "inference_mode": "OpenRouter paid API",
    },
)

# Exact parsing-failure counts from the immediately preceding evaluation
# snapshot, before the parser normalization update.  The evaluation JSON files
# were overwritten by the updated evaluation, so only the preserved exact
# failure counts are compared here.
LEGACY_PARSE_FAILURES = {
    "gpt5": 0,
    "claude45": 34,
    "qwen3_8b": 1870,
    "gpt_oss20b": 71,
}

NAVY = "17365D"
BLUE = "4472C4"
MEDIUM_BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FB"
GREEN = "70AD47"
LIGHT_GREEN = "E2F0D9"
ORANGE = "ED7D31"
LIGHT_ORANGE = "FCE4D6"
RED = "C00000"
LIGHT_RED = "F4CCCC"
GRAY = "7F8C8D"
LIGHT_GRAY = "E7E6E6"
VERY_LIGHT_GRAY = "F3F6F8"
WHITE = "FFFFFF"
BLACK = "000000"
THIN_GRAY = Side(style="thin", color="D9E1F2")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
SECTION_FILL = PatternFill("solid", fgColor=LIGHT_BLUE)
SUBHEADER_FILL = PatternFill("solid", fgColor=BLUE)
GOOD_FILL = PatternFill("solid", fgColor=LIGHT_GREEN)
WARN_FILL = PatternFill("solid", fgColor=LIGHT_ORANGE)
BAD_FILL = PatternFill("solid", fgColor=LIGHT_RED)

PERCENT_HEADERS = {
    "Parsing Failure Rate",
    "Overall Precision",
    "Overall Recall",
    "Overall F1",
    "Pref Precision",
    "Pref Recall",
    "Pref F1",
    "Pref Exact Match",
    "Non-pref Precision",
    "Non-pref Recall",
    "Non-pref F1",
    "Before Failure Rate",
    "After Failure Rate",
    "Failure Reduction Rate",
    "Prediction Equality Rate",
    "Both Pref EM Correct Rate",
    "A-only Pref EM Correct Rate",
    "B-only Pref EM Correct Rate",
    "Neither Pref EM Correct Rate",
    "A Win Rate",
    "Tie Rate",
    "B Win Rate",
    "Consensus Prediction Rate",
}
INTEGER_HEADERS = {
    "N",
    "API Errors",
    "Parsing Failures",
    "Pref Exact Match Count",
    "Pref Total",
    "Total Tokens",
    "Input Tokens",
    "Output Tokens",
    "Reasoning Tokens",
    "Cached Input Tokens",
    "Cache Creation Input Tokens",
    "Population Index",
    "Overall TP",
    "Overall FP",
    "Overall FN",
    "Pref TP",
    "Pref FP",
    "Pref FN",
    "Non-pref TP",
    "Non-pref FP",
    "Non-pref FN",
    "GT Slot Count",
    "Pred Slot Count",
    "Before Parsing Failures",
    "After Parsing Failures",
    "Recovered Failures",
    "Compared Instances",
    "Prediction Equal Count",
    "Both Pref EM Correct",
    "A-only Pref EM Correct",
    "B-only Pref EM Correct",
    "Neither Pref EM Correct",
    "A Wins",
    "Ties",
    "B Wins",
    "Inference Rows",
    "Expected Rows",
    "Reused Rows",
    "New Rows",
    "Unique Example IDs",
    "Unique Example Sub-IDs",
    "Reusable Mix600 Rows",
    "Models Pref EM Correct",
    "Models Parse Failed",
}


@dataclass
class ModelRun:
    key: str
    name: str
    directory: str
    provider: str
    request_model: str
    reasoning: str
    inference_mode: str
    path: Path
    evaluation_path: Path
    inference_path: Path
    report: dict[str, Any]
    run_summary: dict[str, Any]
    reused_rows: int | None
    new_rows: int | None


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--experiment-root",
        type=Path,
        default=DEFAULT_EXPERIMENT_ROOT,
        help="Root containing population/ and vanilla_llm/.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="Destination .xlsx file.",
    )
    parser.add_argument(
        "--skip-instance-sheets",
        action="store_true",
        help="Omit the large per-instance audit sheets.",
    )
    return parser.parse_args()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> Iterable[dict[str, Any]]:
    with path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            value = json.loads(line)
            if not isinstance(value, dict):
                raise ValueError(f"{path}:{line_number} is not a JSON object")
            yield value


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def safe_text(value: Any, *, limit: int = 32760) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    if text[:1] in {"=", "+", "-", "@"}:
        text = "'" + text
    if len(text) > limit:
        return text[: limit - 20] + "\n...[truncated]"
    return text


def load_population(experiment_root: Path) -> list[dict[str, Any]]:
    path = experiment_root / "population" / "population.jsonl"
    if not path.is_file():
        raise FileNotFoundError(f"Missing population file: {path}")
    rows = list(read_jsonl(path))
    rows.sort(key=lambda row: int(row["population_index"]))
    indices = [int(row["population_index"]) for row in rows]
    if indices != list(range(len(rows))):
        raise ValueError("Population indices are not contiguous from zero")
    return rows


def load_runs(
    experiment_root: Path,
    expected_n: int,
) -> list[ModelRun]:
    vanilla_root = experiment_root / "vanilla_llm"
    reuse_path = experiment_root / "reuse_summary.json"
    reuse = read_json(reuse_path) if reuse_path.is_file() else {}
    reuse_models = reuse.get("vanilla_llm", {})
    runs: list[ModelRun] = []
    for spec in MODEL_SPECS:
        path = vanilla_root / str(spec["directory"])
        evaluation_path = path / "evaluation.json"
        inference_path = path / "inference.jsonl"
        if not evaluation_path.is_file():
            raise FileNotFoundError(f"Missing evaluation file: {evaluation_path}")
        if not inference_path.is_file():
            raise FileNotFoundError(f"Missing inference file: {inference_path}")
        report = read_json(evaluation_path)
        actual_n = int(report["overall"]["n"])
        if actual_n != expected_n:
            raise ValueError(
                f"{evaluation_path}: expected {expected_n} rows, got {actual_n}"
            )
        run_summary_path = path / "run_summary.json"
        run_summary = (
            read_json(run_summary_path) if run_summary_path.is_file() else {}
        )
        reuse_record = reuse_models.get(spec["directory"], {})
        reused_rows = reuse_record.get("reused_rows")
        new_rows = report.get("new_batch_rows")
        if new_rows is None and reused_rows is not None:
            new_rows = expected_n - int(reused_rows)
        if reused_rows is None and new_rows is not None:
            reused_rows = expected_n - int(new_rows)
        runs.append(
            ModelRun(
                **spec,
                path=path,
                evaluation_path=evaluation_path,
                inference_path=inference_path,
                report=report,
                run_summary=run_summary,
                reused_rows=int(reused_rows) if reused_rows is not None else None,
                new_rows=int(new_rows) if new_rows is not None else None,
            )
        )
    return runs


def load_metrics_helpers() -> dict[str, Any]:
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from src.evaluation.metrics import (  # noqa: PLC0415
        build_gt_allowed_map,
        build_pred_map,
        counts_slot_and_value_or,
        extract_calls,
        filter_map_by_pref,
        load_pref_list,
        prf_from_counts,
    )

    return {
        "build_gt_allowed_map": build_gt_allowed_map,
        "build_pred_map": build_pred_map,
        "counts_slot_and_value_or": counts_slot_and_value_or,
        "extract_calls": extract_calls,
        "filter_map_by_pref": filter_map_by_pref,
        "load_pref_list": load_pref_list,
        "prf_from_counts": prf_from_counts,
    }


def canonical_slot_map(
    slot_map: Mapping[tuple[str, str], set[str]],
) -> str:
    serializable = [
        [domain, slot, sorted(values)]
        for (domain, slot), values in sorted(slot_map.items())
    ]
    return json.dumps(serializable, ensure_ascii=False, separators=(",", ":"))


def evaluate_instance(
    inference: Mapping[str, Any],
    population: Mapping[str, Any],
    pref_map: Mapping[str, set[str]],
    helpers: Mapping[str, Any],
) -> dict[str, Any]:
    gt_map = helpers["build_gt_allowed_map"](population["reference_ground_truth"])
    pred_map = helpers["build_pred_map"](inference.get("llm_output"))
    calls = helpers["extract_calls"](inference.get("llm_output"))

    def metric(want_pref: bool | None) -> tuple[Any, int, int, int]:
        selected_gt = gt_map
        selected_pred = pred_map
        if want_pref is not None:
            selected_gt = helpers["filter_map_by_pref"](
                gt_map, pref_map, want_pref
            )
            selected_pred = helpers["filter_map_by_pref"](
                pred_map, pref_map, want_pref
            )
        tp, fp, fn = helpers["counts_slot_and_value_or"](
            selected_gt, selected_pred
        )
        return helpers["prf_from_counts"](tp, fp, fn), tp, fp, fn

    overall, overall_tp, overall_fp, overall_fn = metric(None)
    pref, pref_tp, pref_fp, pref_fn = metric(True)
    nonpref, nonpref_tp, nonpref_fp, nonpref_fn = metric(False)
    token_counts = inference.get("token_counts") or {}
    return {
        "population_index": int(population["population_index"]),
        "sample_id": str(population["sample_id"]),
        "example_id": population.get("example_id", ""),
        "example_id_sub": population.get("example_id_sub", ""),
        "turn": population["turn"],
        "difficulty": population["pref_type"],
        "conflict": bool(population.get("conflict", False)),
        "condition": population["condition"],
        "reusable_mix600": bool(population.get("reusable_mix600_example", False)),
        "status": inference.get("status", ""),
        "error": inference.get("error", ""),
        "parse_failed": not bool(calls),
        "parse_reason": "no_calls_extracted" if not calls else "ok",
        "pref_exact_match": 1 if pref_fp == 0 and pref_fn == 0 else 0,
        "overall_precision": overall.precision,
        "overall_recall": overall.recall,
        "overall_f1": overall.f1,
        "pref_precision": pref.precision,
        "pref_recall": pref.recall,
        "pref_f1": pref.f1,
        "nonpref_precision": nonpref.precision,
        "nonpref_recall": nonpref.recall,
        "nonpref_f1": nonpref.f1,
        "overall_tp": overall_tp,
        "overall_fp": overall_fp,
        "overall_fn": overall_fn,
        "pref_tp": pref_tp,
        "pref_fp": pref_fp,
        "pref_fn": pref_fn,
        "nonpref_tp": nonpref_tp,
        "nonpref_fp": nonpref_fp,
        "nonpref_fn": nonpref_fn,
        "gt_slot_count": len(gt_map),
        "pred_slot_count": len(pred_map),
        "ground_truth": safe_text(population["reference_ground_truth"]),
        "parsed_calls": safe_text(calls),
        "parsed_prediction": canonical_slot_map(pred_map),
        "raw_output": safe_text(inference.get("llm_output", "")),
        "reasoning": safe_text(inference.get("reasoning_content", "")),
        "test_utterance": safe_text(
            inference.get("test_utterance", population.get("utterance", ""))
        ),
        "input_tokens": int(token_counts.get("input_tokens", 0) or 0),
        "output_tokens": int(token_counts.get("output_tokens", 0) or 0),
        "reasoning_tokens": int(token_counts.get("reasoning_tokens", 0) or 0),
        "cached_input_tokens": int(
            token_counts.get("cached_input_tokens", 0) or 0
        ),
        "cache_creation_input_tokens": int(
            token_counts.get("cache_creation_input_tokens", 0) or 0
        ),
        "total_tokens": int(token_counts.get("total_tokens", 0) or 0),
        "elapsed_seconds": inference.get("elapsed_seconds"),
        "finish_reason": inference.get("finish_reason", ""),
        "reused_inference": bool(inference.get("reuse_provenance")),
    }


def build_instance_results(
    runs: Sequence[ModelRun],
    population: Sequence[dict[str, Any]],
) -> dict[str, list[dict[str, Any]]]:
    helpers = load_metrics_helpers()
    pref_map = helpers["load_pref_list"](str(ROOT / "config" / "pref_list.json"))
    population_by_index = {
        int(row["population_index"]): row for row in population
    }
    results: dict[str, list[dict[str, Any]]] = {}
    for run in runs:
        run_rows: list[dict[str, Any]] = []
        latest: dict[int, dict[str, Any]] = {}
        attempt_counts: Counter[int] = Counter()
        selected_line: dict[int, int] = {}
        for line_number, inference in enumerate(
            read_jsonl(run.inference_path),
            start=1,
        ):
            index = int(inference["population_index"])
            if index not in population_by_index:
                raise ValueError(f"{run.inference_path}: unknown index {index}")
            attempt_counts[index] += 1
            latest[index] = inference
            selected_line[index] = line_number
        for index, inference in sorted(latest.items()):
            row = evaluate_instance(
                inference,
                population_by_index[index],
                pref_map,
                helpers,
            )
            row.update(
                {
                    "model_key": run.key,
                    "model": run.name,
                    "provider": run.provider,
                    "request_model": run.request_model,
                    "checkpoint_attempt_count": attempt_counts[index],
                    "selected_checkpoint_line": selected_line[index],
                }
            )
            run_rows.append(row)
        run_rows.sort(key=lambda row: row["population_index"])
        if len(run_rows) != len(population):
            raise ValueError(
                f"{run.inference_path}: expected {len(population)} rows, "
                f"got {len(run_rows)}"
            )
        expected_parse_failures = int(
            run.report["overall"]["parsing_failures"]
        )
        actual_parse_failures = sum(
            bool(row["parse_failed"]) for row in run_rows
        )
        if actual_parse_failures != expected_parse_failures:
            raise ValueError(
                f"{run.inference_path}: selected latest rows produced "
                f"{actual_parse_failures} parse failures; evaluation.json "
                f"reports {expected_parse_failures}"
            )
        expected_pref_em = int(
            run.report["overall"]["pref"]["exact_match_count"]
        )
        actual_pref_em = sum(row["pref_exact_match"] for row in run_rows)
        if actual_pref_em != expected_pref_em:
            raise ValueError(
                f"{run.inference_path}: selected latest rows produced "
                f"{actual_pref_em} preference exact matches; evaluation.json "
                f"reports {expected_pref_em}"
            )
        results[run.key] = run_rows
    return results


def metric_columns(metric_row: Mapping[str, Any]) -> list[Any]:
    return [
        metric_row["n"],
        metric_row["api_errors"],
        metric_row["parsing_failures"],
        metric_row["parsing_failure_rate"],
        metric_row["overall"]["precision"],
        metric_row["overall"]["recall"],
        metric_row["overall"]["f1"],
        metric_row["pref"]["precision"],
        metric_row["pref"]["recall"],
        metric_row["pref"]["f1"],
        metric_row["pref"]["exact_match_rate"],
        metric_row["pref"]["exact_match_count"],
        metric_row["pref"]["total"],
        metric_row["nonpref"]["precision"],
        metric_row["nonpref"]["recall"],
        metric_row["nonpref"]["f1"],
    ]


METRIC_HEADERS = [
    "N",
    "API Errors",
    "Parsing Failures",
    "Parsing Failure Rate",
    "Overall Precision",
    "Overall Recall",
    "Overall F1",
    "Pref Precision",
    "Pref Recall",
    "Pref F1",
    "Pref Exact Match",
    "Pref Exact Match Count",
    "Pref Total",
    "Non-pref Precision",
    "Non-pref Recall",
    "Non-pref F1",
]


def dense_ranks(values: Sequence[float], *, higher_is_better: bool) -> list[int]:
    unique = sorted(set(values), reverse=higher_is_better)
    lookup = {value: rank for rank, value in enumerate(unique, start=1)}
    return [lookup[value] for value in values]


def title(
    ws: Any,
    text: str,
    subtitle: str,
    *,
    last_column: int,
) -> None:
    end = get_column_letter(last_column)
    ws.merge_cells(f"A1:{end}1")
    ws["A1"] = text
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(f"A2:{end}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, italic=True, color="44546A")
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 28


def section(ws: Any, row: int, text: str, *, last_column: int) -> None:
    ws.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=last_column,
    )
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(size=12, bold=True, color=NAVY)
    cell.fill = SECTION_FILL
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 21


def header(ws: Any, row: int, values: Sequence[str]) -> None:
    for column, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=column, value=value)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = Border(bottom=THIN_GRAY)
    ws.row_dimensions[row].height = 38


def add_table(
    ws: Any,
    *,
    name: str,
    start_row: int,
    end_row: int,
    end_column: int,
) -> None:
    if end_row <= start_row:
        return
    ref = f"A{start_row}:{get_column_letter(end_column)}{end_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def format_by_header(
    ws: Any,
    headers: Sequence[str],
    *,
    first_data_row: int,
    last_data_row: int,
) -> None:
    for column, name in enumerate(headers, start=1):
        for row in range(first_data_row, last_data_row + 1):
            cell = ws.cell(row=row, column=column)
            if name in PERCENT_HEADERS or any(
                token in name
                for token in (
                    " F1",
                    " Precision",
                    " Recall",
                    " Exact Match",
                    "Failure Rate",
                    " Win Rate",
                    "Tie Rate",
                    "Equality Rate",
                )
            ):
                cell.number_format = "0.00%"
            elif name in INTEGER_HEADERS or name.endswith(" Rank"):
                cell.number_format = "#,##0"
            elif "Tokens / Instance" in name:
                cell.number_format = "#,##0.0"
            elif name in {"Cost (USD)", "Cost / 1K Instances (USD)"}:
                cell.number_format = '$0.0000'
            elif name in {"Elapsed Seconds", "Mean Overall F1", "F1 Range"}:
                cell.number_format = "0.0000"


def apply_metric_color_scales(
    ws: Any,
    headers: Sequence[str],
    *,
    first_data_row: int,
    last_data_row: int,
) -> None:
    for column, name in enumerate(headers, start=1):
        if not any(
            marker in name
            for marker in (
                " F1",
                " Exact Match",
                " Precision",
                " Recall",
            )
        ):
            continue
        letter = get_column_letter(column)
        ws.conditional_formatting.add(
            f"{letter}{first_data_row}:{letter}{last_data_row}",
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )


def finish_sheet(
    ws: Any,
    *,
    freeze: str = "A5",
    landscape: bool = True,
) -> None:
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = ws.dimensions
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5


def set_widths(
    ws: Any,
    headers: Sequence[str],
    *,
    custom: Mapping[str, float] | None = None,
) -> None:
    custom = custom or {}
    for column, name in enumerate(headers, start=1):
        if name in custom:
            width = custom[name]
        elif any(token in name for token in ("Path", "SHA256", "Output")):
            width = 42
        elif name in {"Model", "Request Model", "Inference Mode"}:
            width = 24
        elif name in {
            "Example ID",
            "Example Sub-ID",
            "Ground Truth",
            "Parsed Calls",
            "Parsed Prediction",
            "Reasoning",
            "Test Utterance",
            "Error",
            "Notes",
        }:
            width = 34
        else:
            width = max(11, min(24, len(name) + 2))
        ws.column_dimensions[get_column_letter(column)].width = width


def write_standard_sheet(
    wb: Workbook,
    *,
    name: str,
    title_text: str,
    subtitle: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    table_name: str,
    freeze: str = "A5",
    custom_widths: Mapping[str, float] | None = None,
    color_scales: bool = True,
) -> Any:
    ws = wb.create_sheet(name)
    title(ws, title_text, subtitle, last_column=len(headers))
    header(ws, 4, headers)
    for values in rows:
        ws.append(list(values))
    last_row = 4 + len(rows)
    add_table(
        ws,
        name=table_name,
        start_row=4,
        end_row=last_row,
        end_column=len(headers),
    )
    format_by_header(
        ws,
        headers,
        first_data_row=5,
        last_data_row=last_row,
    )
    if color_scales and rows:
        apply_metric_color_scales(
            ws,
            headers,
            first_data_row=5,
            last_data_row=last_row,
        )
    set_widths(ws, headers, custom=custom_widths)
    finish_sheet(ws, freeze=freeze)
    return ws


def overall_rows(runs: Sequence[ModelRun]) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "Model",
        "Provider",
        "Request Model",
        "Reasoning",
        "Inference Mode",
        *METRIC_HEADERS,
        "Overall F1 Rank",
        "Pref F1 Rank",
        "Pref EM Rank",
        "Non-pref F1 Rank",
        "Total Tokens",
        "Input Tokens",
        "Output Tokens",
        "Reasoning Tokens",
        "Cached Input Tokens",
        "Cache Creation Input Tokens",
        "Tokens / Instance",
        "Cost (USD)",
        "Cost / 1K Instances (USD)",
        "Reused Rows",
        "New Rows",
        "Evaluation Generated At",
        "Inference SHA256",
        "Evaluation Path",
    ]
    reports = [run.report["overall"] for run in runs]
    rank_sets = {
        "overall": dense_ranks(
            [row["overall"]["f1"] for row in reports],
            higher_is_better=True,
        ),
        "pref": dense_ranks(
            [row["pref"]["f1"] for row in reports],
            higher_is_better=True,
        ),
        "em": dense_ranks(
            [row["pref"]["exact_match_rate"] for row in reports],
            higher_is_better=True,
        ),
        "nonpref": dense_ranks(
            [row["nonpref"]["f1"] for row in reports],
            higher_is_better=True,
        ),
    }
    rows: list[list[Any]] = []
    for index, run in enumerate(runs):
        report = run.report
        overall = report["overall"]
        usage = report.get("token_usage", {})
        n = int(overall["n"])
        total_tokens = int(usage.get("total_tokens", 0) or 0)
        cost = run.run_summary.get("token_usage", {}).get("cost_usd")
        rows.append(
            [
                run.name,
                run.provider,
                run.request_model,
                run.reasoning,
                run.inference_mode,
                *metric_columns(overall),
                rank_sets["overall"][index],
                rank_sets["pref"][index],
                rank_sets["em"][index],
                rank_sets["nonpref"][index],
                total_tokens,
                int(usage.get("input_tokens", 0) or 0),
                int(usage.get("output_tokens", 0) or 0),
                int(usage.get("reasoning_tokens", 0) or 0),
                int(usage.get("cached_input_tokens", 0) or 0),
                int(usage.get("cache_creation_input_tokens", 0) or 0),
                total_tokens / n if n else 0,
                cost,
                (float(cost) / n * 1000) if cost is not None and n else None,
                run.reused_rows,
                run.new_rows,
                report.get("generated_at", ""),
                report.get("inference_sha256", ""),
                str(run.evaluation_path),
            ]
        )
    return headers, rows


def condition_rows(runs: Sequence[ModelRun]) -> tuple[list[str], list[list[Any]]]:
    headers = ["Model", "Turn", "Difficulty", "Condition", *METRIC_HEADERS]
    rows: list[list[Any]] = []
    for run in runs:
        for condition in CONDITION_ORDER:
            metric = run.report["conditions"][condition]
            turn, difficulty = condition.split("_", 1)
            rows.append(
                [run.name, turn, difficulty, condition, *metric_columns(metric)]
            )
    return headers, rows


def split_rows(runs: Sequence[ModelRun]) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "Model",
        "Turn",
        "Difficulty",
        "Conflict Group",
        "Condition",
        *METRIC_HEADERS,
    ]
    rows: list[list[Any]] = []
    for run in runs:
        for condition in CONDITION_ORDER:
            turn, difficulty = condition.split("_", 1)
            for conflict in CONFLICT_ORDER:
                metric = run.report["condition_conflict"][condition][conflict]
                values = metric_columns(metric)
                if metric["n"] == 0:
                    values = values[:4] + [None] * (len(values) - 4)
                rows.append(
                    [
                        run.name,
                        turn,
                        difficulty,
                        conflict,
                        condition,
                        *values,
                    ]
                )
    return headers, rows


def ranking_rows(runs: Sequence[ModelRun]) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "Turn",
        "Difficulty",
        "Conflict Group",
        "N",
        "Model",
        "Overall F1",
        "Overall F1 Rank",
        "Pref F1",
        "Pref F1 Rank",
        "Pref Exact Match",
        "Pref EM Rank",
        "Non-pref F1",
        "Non-pref F1 Rank",
        "Parsing Failure Rate",
        "Parsing Quality Rank",
    ]
    rows: list[list[Any]] = []
    for turn in TURN_ORDER:
        for difficulty in DIFFICULTY_ORDER:
            condition = f"{turn}_{difficulty}"
            for conflict in CONFLICT_ORDER:
                metrics = [
                    run.report["condition_conflict"][condition][conflict]
                    for run in runs
                ]
                if not metrics or int(metrics[0]["n"]) == 0:
                    continue
                overall_rank = dense_ranks(
                    [row["overall"]["f1"] for row in metrics],
                    higher_is_better=True,
                )
                pref_rank = dense_ranks(
                    [row["pref"]["f1"] for row in metrics],
                    higher_is_better=True,
                )
                em_rank = dense_ranks(
                    [row["pref"]["exact_match_rate"] for row in metrics],
                    higher_is_better=True,
                )
                nonpref_rank = dense_ranks(
                    [row["nonpref"]["f1"] for row in metrics],
                    higher_is_better=True,
                )
                parse_rank = dense_ranks(
                    [row["parsing_failure_rate"] for row in metrics],
                    higher_is_better=False,
                )
                for i, run in enumerate(runs):
                    metric = metrics[i]
                    rows.append(
                        [
                            turn,
                            difficulty,
                            conflict,
                            metric["n"],
                            run.name,
                            metric["overall"]["f1"],
                            overall_rank[i],
                            metric["pref"]["f1"],
                            pref_rank[i],
                            metric["pref"]["exact_match_rate"],
                            em_rank[i],
                            metric["nonpref"]["f1"],
                            nonpref_rank[i],
                            metric["parsing_failure_rate"],
                            parse_rank[i],
                        ]
                    )
    return headers, rows


def conflict_gap_rows(
    runs: Sequence[ModelRun],
) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "Model",
        "Turn",
        "Difficulty",
        "Non-conflict N",
        "Conflict N",
        "Non-conflict Overall F1",
        "Conflict Overall F1",
        "Conflict − Non-conflict Overall F1",
        "Non-conflict Pref F1",
        "Conflict Pref F1",
        "Conflict − Non-conflict Pref F1",
        "Non-conflict Pref Exact Match",
        "Conflict Pref Exact Match",
        "Conflict − Non-conflict Pref EM",
        "Non-conflict Non-pref F1",
        "Conflict Non-pref F1",
        "Conflict − Non-conflict Non-pref F1",
        "Non-conflict Parsing Failure Rate",
        "Conflict Parsing Failure Rate",
        "Conflict − Non-conflict Failure Rate",
    ]
    rows: list[list[Any]] = []
    for run in runs:
        for condition in CONDITION_ORDER:
            turn, difficulty = condition.split("_", 1)
            nonconf = run.report["condition_conflict"][condition]["non_conflict"]
            conflict = run.report["condition_conflict"][condition]["conflict"]
            if not conflict["n"]:
                gaps = [None] * 15
                rows.append(
                    [
                        run.name,
                        turn,
                        difficulty,
                        nonconf["n"],
                        0,
                        *gaps,
                    ]
                )
                continue
            rows.append(
                [
                    run.name,
                    turn,
                    difficulty,
                    nonconf["n"],
                    conflict["n"],
                    nonconf["overall"]["f1"],
                    conflict["overall"]["f1"],
                    conflict["overall"]["f1"] - nonconf["overall"]["f1"],
                    nonconf["pref"]["f1"],
                    conflict["pref"]["f1"],
                    conflict["pref"]["f1"] - nonconf["pref"]["f1"],
                    nonconf["pref"]["exact_match_rate"],
                    conflict["pref"]["exact_match_rate"],
                    conflict["pref"]["exact_match_rate"]
                    - nonconf["pref"]["exact_match_rate"],
                    nonconf["nonpref"]["f1"],
                    conflict["nonpref"]["f1"],
                    conflict["nonpref"]["f1"] - nonconf["nonpref"]["f1"],
                    nonconf["parsing_failure_rate"],
                    conflict["parsing_failure_rate"],
                    conflict["parsing_failure_rate"]
                    - nonconf["parsing_failure_rate"],
                ]
            )
    return headers, rows


def token_rows(runs: Sequence[ModelRun]) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "Model",
        "Provider",
        "N",
        "Total Tokens",
        "Input Tokens",
        "Output Tokens",
        "Reasoning Tokens",
        "Cached Input Tokens",
        "Cache Creation Input Tokens",
        "Total Tokens / Instance",
        "Input Tokens / Instance",
        "Output Tokens / Instance",
        "Reasoning Tokens / Instance",
        "Cached Input Tokens / Instance",
        "Cost (USD)",
        "Cost / 1K Instances (USD)",
        "Elapsed Seconds",
        "Notes",
    ]
    rows: list[list[Any]] = []
    for run in runs:
        usage = run.report.get("token_usage", {})
        n = int(run.report["overall"]["n"])
        total = int(usage.get("total_tokens", 0) or 0)
        input_tokens = int(usage.get("input_tokens", 0) or 0)
        output_tokens = int(usage.get("output_tokens", 0) or 0)
        reasoning_tokens = int(usage.get("reasoning_tokens", 0) or 0)
        cached = int(usage.get("cached_input_tokens", 0) or 0)
        cache_creation = int(
            usage.get("cache_creation_input_tokens", 0) or 0
        )
        run_usage = run.run_summary.get("token_usage", {})
        cost = run_usage.get("cost_usd")
        note = (
            "Exact cost preserved in run_summary.json"
            if cost is not None
            else "Cost not preserved in evaluation artifacts"
        )
        rows.append(
            [
                run.name,
                run.provider,
                n,
                total,
                input_tokens,
                output_tokens,
                reasoning_tokens,
                cached,
                cache_creation,
                total / n,
                input_tokens / n,
                output_tokens / n,
                reasoning_tokens / n,
                cached / n,
                cost,
                (float(cost) / n * 1000) if cost is not None else None,
                run.run_summary.get("elapsed_seconds"),
                note,
            ]
        )
    return headers, rows


def parser_impact_rows(
    runs: Sequence[ModelRun],
) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "Model",
        "N",
        "Before Parsing Failures",
        "Before Failure Rate",
        "After Parsing Failures",
        "After Failure Rate",
        "Recovered Failures",
        "Failure Reduction Rate",
        "Current Overall F1",
        "Current Pref F1",
        "Current Pref Exact Match",
        "Parser Update Coverage",
    ]
    rows: list[list[Any]] = []
    for run in runs:
        current = run.report["overall"]
        n = int(current["n"])
        before = LEGACY_PARSE_FAILURES[run.key]
        after = int(current["parsing_failures"])
        reduction = ((before - after) / before) if before else None
        rows.append(
            [
                run.name,
                n,
                before,
                before / n,
                after,
                after / n,
                before - after,
                reduction,
                current["overall"]["f1"],
                current["pref"]["f1"],
                current["pref"]["exact_match_rate"],
                (
                    "Markdown bold functions; quoted/braced Qwen calls; "
                    "comma-wrapped calls; quoted slots; one missing ')' repair"
                ),
            ]
        )
    return headers, rows


def population_composition_rows(
    population: Sequence[dict[str, Any]],
) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "Turn",
        "Difficulty",
        "Conflict Group",
        "N",
        "Share of Population",
        "Unique Example IDs",
        "Unique Example Sub-IDs",
        "Reusable Mix600 Rows",
        "New/Conflict Rows",
    ]
    rows: list[list[Any]] = []
    total = len(population)
    for turn in TURN_ORDER:
        for difficulty in DIFFICULTY_ORDER:
            for label, conflict in (
                ("non_conflict", False),
                ("conflict", True),
            ):
                subset = [
                    row
                    for row in population
                    if row["turn"] == turn
                    and row["pref_type"] == difficulty
                    and bool(row.get("conflict")) is conflict
                ]
                reusable = sum(
                    bool(row.get("reusable_mix600_example")) for row in subset
                )
                rows.append(
                    [
                        turn,
                        difficulty,
                        label,
                        len(subset),
                        len(subset) / total if total else 0,
                        len({row["example_id"] for row in subset}),
                        len({row["example_id_sub"] for row in subset}),
                        reusable,
                        len(subset) - reusable,
                    ]
                )
    return headers, rows


def run_inventory_rows(
    experiment_root: Path,
    included_runs: Sequence[ModelRun],
    expected_n: int,
) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "Directory",
        "Included in Workbook",
        "Evaluation Present",
        "Inference Present",
        "Raw Checkpoint Rows",
        "Unique Inference Rows",
        "Duplicate/Retry Rows",
        "Expected Rows",
        "Completion Rate",
        "Status",
        "Evaluation Path",
        "Inference Path",
    ]
    included = {run.directory for run in included_runs}
    rows: list[list[Any]] = []
    vanilla_root = experiment_root / "vanilla_llm"
    for path in sorted(p for p in vanilla_root.iterdir() if p.is_dir()):
        evaluation = path / "evaluation.json"
        inference = path / "inference.jsonl"
        raw_rows = 0
        unique_indices: set[int] = set()
        if inference.is_file():
            for record in read_jsonl(inference):
                raw_rows += 1
                unique_indices.add(int(record["population_index"]))
        inference_rows = len(unique_indices)
        duplicate_rows = raw_rows - inference_rows
        completion = inference_rows / expected_n if expected_n else 0
        status = (
            "Complete / included"
            if path.name in included
            else (
                "Complete / not selected"
                if inference_rows == expected_n and evaluation.is_file()
                else "Incomplete / excluded"
            )
        )
        rows.append(
            [
                path.name,
                "yes" if path.name in included else "no",
                "yes" if evaluation.is_file() else "no",
                "yes" if inference.is_file() else "no",
                raw_rows,
                inference_rows,
                duplicate_rows,
                expected_n,
                completion,
                status,
                str(evaluation) if evaluation.is_file() else "",
                str(inference) if inference.is_file() else "",
            ]
        )
    return headers, rows


def matrix_sheet(
    wb: Workbook,
    runs: Sequence[ModelRun],
    *,
    turn: str,
) -> None:
    ws = wb.create_sheet(f"{turn.title()} Matrix")
    columns = [
        "Easy NC",
        "Easy Conflict",
        "Medium NC",
        "Medium Conflict",
        "Hard NC",
        "Hard Conflict",
    ]
    title(
        ws,
        f"{turn.title()}-turn Performance Matrix",
        "Blank cells mean the split has N=0 (easy conflict was excluded).",
        last_column=1 + len(columns),
    )
    metrics = (
        ("Overall F1", ("overall", "f1")),
        ("Pref F1", ("pref", "f1")),
        ("Pref Exact Match", ("pref", "exact_match_rate")),
        ("Non-pref F1", ("nonpref", "f1")),
        ("Parsing Failure Rate", ("parsing_failure_rate",)),
    )
    row = 4
    for metric_name, path in metrics:
        section(
            ws,
            row,
            metric_name,
            last_column=1 + len(columns),
        )
        header(ws, row + 1, ["Model", *columns])
        start_data = row + 2
        for run in runs:
            values: list[Any] = [run.name]
            for difficulty in DIFFICULTY_ORDER:
                condition = f"{turn}_{difficulty}"
                for conflict in CONFLICT_ORDER:
                    metric = run.report["condition_conflict"][condition][conflict]
                    if metric["n"] == 0:
                        values.append(None)
                        continue
                    value: Any = metric
                    for key in path:
                        value = value[key]
                    values.append(value)
            ws.append(values)
        end_data = start_data + len(runs) - 1
        for data_row in range(start_data, end_data + 1):
            for column in range(2, len(columns) + 2):
                ws.cell(data_row, column).number_format = "0.00%"
        for column in range(2, len(columns) + 2):
            letter = get_column_letter(column)
            ws.conditional_formatting.add(
                f"{letter}{start_data}:{letter}{end_data}",
                ColorScaleRule(
                    start_type="min",
                    start_color="F8696B",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="max",
                    end_color="63BE7B",
                ),
            )
        row = end_data + 2
    ws.column_dimensions["A"].width = 24
    for column in range(2, len(columns) + 2):
        ws.column_dimensions[get_column_letter(column)].width = 16
    finish_sheet(ws, freeze="B6")


def readme_sheet(
    wb: Workbook,
    *,
    experiment_root: Path,
    output_path: Path,
    runs: Sequence[ModelRun],
    population: Sequence[dict[str, Any]],
) -> None:
    ws = wb.active
    ws.title = "README"
    title(
        ws,
        "Experiment8 Vanilla LLM Detailed Results",
        "MPT_v2_0725 · hint query · easy-conflict excluded · current parser",
        last_column=5,
    )
    section(ws, 4, "Experiment scope", last_column=5)
    facts = [
        ("Dataset", str(ROOT / "data" / "MPT_v2_0725.json")),
        (
            "Dataset SHA256",
            (
                read_json(experiment_root / "vanilla_api_batch_resume" / "resume_summary.json").get(
                    "dataset_sha256", ""
                )
                if (
                    experiment_root
                    / "vanilla_api_batch_resume"
                    / "resume_summary.json"
                ).is_file()
                else ""
            ),
        ),
        ("Population file", str(experiment_root / "population" / "population.jsonl")),
        ("Population size", len(population)),
        ("Method", "vanilla_llm"),
        ("Query", "hint"),
        ("Context type", "diag-apilist"),
        ("Easy conflict excluded", "yes"),
        (
            "Conflict query policy",
            (
                "For conflict cases, queries are generated only for the "
                "majority-preference group; no query is attached to the "
                "minority group."
            ),
        ),
        ("Included models", ", ".join(run.name for run in runs)),
        (
            "Excluded/incomplete run",
            "Gemma 4 12B: inference exists but is incomplete, so it is excluded.",
        ),
        ("Workbook generated (UTC)", datetime.now(timezone.utc).isoformat()),
        ("Workbook path", str(output_path)),
    ]
    row = 5
    for key, value in facts:
        ws.cell(row=row, column=1, value=key).font = Font(bold=True, color=NAVY)
        ws.cell(row=row, column=2, value=safe_text(value))
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
        row += 1

    row += 1
    section(ws, row, "Workbook sheet guide", last_column=5)
    row += 1
    header(ws, row, ["Sheet", "Purpose", "Unit", "Filterable", "Notes"])
    guides = [
        ("README", "Scope, definitions, caveats, provenance", "workbook", "no", ""),
        ("Dashboard", "Headline comparison and charts", "model", "no", ""),
        ("Overall Summary", "All overall metrics, ranks, tokens, settings", "model", "yes", "4 included models"),
        ("Condition Summary", "Single/multi × easy/medium/hard", "model-condition", "yes", "24 rows"),
        ("Split Detail", "Condition split by conflict", "model-split", "yes", "48 rows including zero-N easy-conflict"),
        ("Single Matrix", "Compact single-turn metric matrices", "model-split", "no", ""),
        ("Multi Matrix", "Compact multi-turn metric matrices", "model-split", "no", ""),
        ("Rankings", "Per-split dense ranks for every model", "model-split", "yes", "Only nonempty splits"),
        ("Conflict Gaps", "Conflict minus non-conflict deltas", "model-condition", "yes", ""),
        ("Token Usage", "Token totals, per-instance usage, recorded cost", "model", "yes", "Cost only where preserved"),
        ("Parser Impact", "Exact failure-count change after parser update", "model", "yes", ""),
        ("Dataset Composition", "Population counts and reuse composition", "split", "yes", ""),
        ("Pairwise Agreement", "Pairwise prediction/correctness agreement", "model pair", "yes", ""),
        ("Run Inventory", "Complete and incomplete run artifacts", "run directory", "yes", ""),
        ("Instance Agreement", "Cross-model comparison for each sample", "sample", "yes", "4,695 rows"),
        ("Instance Results", "Auditable per-model, per-sample metrics and outputs", "model-sample", "yes", "18,780 rows"),
        ("Parse Failures", "All remaining current parse failures with raw output", "failed model-sample", "yes", ""),
    ]
    for values in guides:
        ws.append(values)
    guide_end = row + len(guides)
    add_table(
        ws,
        name="SheetGuide",
        start_row=row,
        end_row=guide_end,
        end_column=5,
    )

    row = guide_end + 2
    section(ws, row, "Metric definitions", last_column=5)
    row += 1
    definitions = [
        (
            "Overall Precision / Recall / F1",
            "Micro scores over (domain, slot) pairs. A prediction is correct when any predicted value intersects the OR-allowed ground-truth values.",
        ),
        (
            "Pref Precision / Recall / F1",
            "The same micro metric after retaining only preference slots listed in config/pref_list.json.",
        ),
        (
            "Pref Exact Match",
            "Fraction of instances with zero FP and zero FN among preference slots.",
        ),
        (
            "Non-pref Precision / Recall / F1",
            "The same micro metric after retaining only non-preference slots.",
        ),
        (
            "Parsing Failure",
            "No function call could be extracted from llm_output under the current evaluation parser.",
        ),
        (
            "Conflict Group",
            "The population's explicit conflict flag, derived from conflict_* dataset subsets.",
        ),
        (
            "Ranks",
            "Dense ranks within the same population/split; larger performance metrics are better, while lower parsing-failure rate is better.",
        ),
    ]
    for name, definition in definitions:
        ws.cell(row=row, column=1, value=name).font = Font(bold=True, color=NAVY)
        ws.cell(row=row, column=2, value=definition)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = 36
        row += 1

    row += 1
    section(ws, row, "Interpretation caveats", last_column=5)
    row += 1
    caveats = [
        "Easy-conflict samples are intentionally excluded; zero-N easy-conflict cells are blank rather than reported as zero performance.",
        "Micro-F1 aggregates slot decisions, while Pref Exact Match is instance-level. They answer different questions and should be read together.",
        "GPT-5 and Claude costs are not present in the final evaluation artifacts, so the workbook does not invent them. OpenRouter costs are included where run_summary.json preserves exact values.",
        "Parser Impact compares exact failure counts only. The overwritten pre-update aggregate metric files were not preserved, so pre-update F1 values are not fabricated.",
        "The per-instance sheet omits the full model prompt to keep the workbook practical; it preserves the query utterance, ground truth, parsed calls, raw output, reasoning, token counts, and all scoring counts.",
    ]
    for caveat in caveats:
        ws.cell(row=row, column=1, value="•")
        ws.cell(row=row, column=2, value=caveat)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = 34
        row += 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 34
    finish_sheet(ws, freeze="A5", landscape=False)
    ws.auto_filter.ref = None


def dashboard_sheet(wb: Workbook, runs: Sequence[ModelRun]) -> None:
    ws = wb.create_sheet("Dashboard")
    title(
        ws,
        "Vanilla LLM Performance Dashboard",
        "All models evaluated on the identical 4,695-instance population.",
        last_column=14,
    )
    overall = [run.report["overall"] for run in runs]
    best_f1 = max(range(len(runs)), key=lambda i: overall[i]["overall"]["f1"])
    best_em = max(
        range(len(runs)),
        key=lambda i: overall[i]["pref"]["exact_match_rate"],
    )
    hard_conflict = [
        run.report["condition_conflict"]["multi_hard"]["conflict"]["overall"][
            "f1"
        ]
        for run in runs
    ]
    best_hard_conflict = max(range(len(runs)), key=lambda i: hard_conflict[i])

    section(ws, 4, "Headline findings", last_column=14)
    findings = [
        (
            "Best overall F1",
            runs[best_f1].name,
            overall[best_f1]["overall"]["f1"],
        ),
        (
            "Best preference exact match",
            runs[best_em].name,
            overall[best_em]["pref"]["exact_match_rate"],
        ),
        (
            "Best multi-hard conflict overall F1",
            runs[best_hard_conflict].name,
            hard_conflict[best_hard_conflict],
        ),
        (
            "Lowest parsing failures",
            ", ".join(
                run.name
                for run in runs
                if run.report["overall"]["parsing_failures"]
                == min(
                    item.report["overall"]["parsing_failures"] for item in runs
                )
            ),
            min(run.report["overall"]["parsing_failure_rate"] for run in runs),
        ),
    ]
    header(ws, 5, ["Finding", "Model", "Value"])
    for finding in findings:
        ws.append(finding)
    for row in range(6, 6 + len(findings)):
        ws.cell(row, 3).number_format = "0.00%"
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 16

    row = 11
    section(ws, row, "Overall model comparison", last_column=14)
    headers = [
        "Model",
        "Overall F1",
        "Pref F1",
        "Pref Exact Match",
        "Non-pref F1",
        "Parsing Failure Rate",
        "Total Tokens",
    ]
    header(ws, row + 1, headers)
    for run in runs:
        metric = run.report["overall"]
        ws.append(
            [
                run.name,
                metric["overall"]["f1"],
                metric["pref"]["f1"],
                metric["pref"]["exact_match_rate"],
                metric["nonpref"]["f1"],
                metric["parsing_failure_rate"],
                run.report["token_usage"]["total_tokens"],
            ]
        )
    data_start = row + 2
    data_end = data_start + len(runs) - 1
    for r in range(data_start, data_end + 1):
        for c in range(2, 7):
            ws.cell(r, c).number_format = "0.00%"
        ws.cell(r, 7).number_format = "#,##0"
    add_table(
        ws,
        name="DashboardOverall",
        start_row=row + 1,
        end_row=data_end,
        end_column=len(headers),
    )

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Overall / Preference / Non-preference Performance"
    chart.y_axis.title = "Score"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 1
    chart.x_axis.title = "Model"
    chart.height = 8
    chart.width = 16
    data = Reference(
        ws,
        min_col=2,
        max_col=5,
        min_row=row + 1,
        max_row=data_end,
    )
    categories = Reference(
        ws,
        min_col=1,
        min_row=data_start,
        max_row=data_end,
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    ws.add_chart(chart, "I5")

    parse_chart = BarChart()
    parse_chart.type = "bar"
    parse_chart.style = 12
    parse_chart.title = "Parsing Failure Rate"
    parse_chart.x_axis.title = "Failure rate"
    parse_chart.x_axis.scaling.min = 0
    parse_chart.height = 7
    parse_chart.width = 16
    parse_data = Reference(
        ws,
        min_col=6,
        min_row=row + 1,
        max_row=data_end,
    )
    parse_chart.add_data(parse_data, titles_from_data=True)
    parse_chart.set_categories(categories)
    ws.add_chart(parse_chart, "I20")
    finish_sheet(ws, freeze="A5")
    ws.auto_filter.ref = None


def pairwise_rows(
    runs: Sequence[ModelRun],
    instance_results: Mapping[str, Sequence[dict[str, Any]]],
) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "Model A",
        "Model B",
        "Compared Instances",
        "Prediction Equal Count",
        "Prediction Equality Rate",
        "Both Pref EM Correct",
        "Both Pref EM Correct Rate",
        "A-only Pref EM Correct",
        "A-only Pref EM Correct Rate",
        "B-only Pref EM Correct",
        "B-only Pref EM Correct Rate",
        "Neither Pref EM Correct",
        "Neither Pref EM Correct Rate",
        "A Wins",
        "A Win Rate",
        "Ties",
        "Tie Rate",
        "B Wins",
        "B Win Rate",
        "Mean Overall F1 A",
        "Mean Overall F1 B",
        "Mean F1 Difference A − B",
    ]
    rows: list[list[Any]] = []
    run_by_key = {run.key: run for run in runs}
    for key_a, key_b in combinations((run.key for run in runs), 2):
        a_rows = instance_results[key_a]
        b_rows = instance_results[key_b]
        n = len(a_rows)
        equal = both = a_only = b_only = neither = 0
        a_wins = ties = b_wins = 0
        sum_a = sum_b = 0.0
        for a, b in zip(a_rows, b_rows):
            if a["population_index"] != b["population_index"]:
                raise ValueError("Instance rows are not aligned")
            equal += a["parsed_prediction"] == b["parsed_prediction"]
            a_em = bool(a["pref_exact_match"])
            b_em = bool(b["pref_exact_match"])
            both += a_em and b_em
            a_only += a_em and not b_em
            b_only += b_em and not a_em
            neither += not a_em and not b_em
            a_f1 = float(a["overall_f1"])
            b_f1 = float(b["overall_f1"])
            sum_a += a_f1
            sum_b += b_f1
            if math.isclose(a_f1, b_f1, abs_tol=1e-12):
                ties += 1
            elif a_f1 > b_f1:
                a_wins += 1
            else:
                b_wins += 1
        rows.append(
            [
                run_by_key[key_a].name,
                run_by_key[key_b].name,
                n,
                equal,
                equal / n,
                both,
                both / n,
                a_only,
                a_only / n,
                b_only,
                b_only / n,
                neither,
                neither / n,
                a_wins,
                a_wins / n,
                ties,
                ties / n,
                b_wins,
                b_wins / n,
                sum_a / n,
                sum_b / n,
                (sum_a - sum_b) / n,
            ]
        )
    return headers, rows


def instance_agreement_rows(
    runs: Sequence[ModelRun],
    population: Sequence[dict[str, Any]],
    instance_results: Mapping[str, Sequence[dict[str, Any]]],
) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "Population Index",
        "Sample ID",
        "Example ID",
        "Example Sub-ID",
        "Turn",
        "Difficulty",
        "Conflict Group",
        "Condition",
        "Ground Truth",
    ]
    for run in runs:
        headers.extend(
            [
                f"{run.name} Overall F1",
                f"{run.name} Pref Exact Match",
                f"{run.name} Parse Failed",
                f"{run.name} Parsed Prediction",
            ]
        )
    headers.extend(
        [
            "Models Pref EM Correct",
            "Models Parse Failed",
            "Mean Overall F1",
            "Minimum Overall F1",
            "Maximum Overall F1",
            "F1 Range",
            "All Parsed Predictions Equal",
        ]
    )
    rows: list[list[Any]] = []
    for index, pop in enumerate(population):
        model_rows = [instance_results[run.key][index] for run in runs]
        f1_values = [float(row["overall_f1"]) for row in model_rows]
        values: list[Any] = [
            pop["population_index"],
            pop["sample_id"],
            pop.get("example_id", ""),
            pop.get("example_id_sub", ""),
            pop["turn"],
            pop["pref_type"],
            "conflict" if pop.get("conflict") else "non_conflict",
            pop["condition"],
            safe_text(pop["reference_ground_truth"]),
        ]
        for model_row in model_rows:
            values.extend(
                [
                    model_row["overall_f1"],
                    model_row["pref_exact_match"],
                    "yes" if model_row["parse_failed"] else "no",
                    model_row["parsed_prediction"],
                ]
            )
        values.extend(
            [
                sum(row["pref_exact_match"] for row in model_rows),
                sum(row["parse_failed"] for row in model_rows),
                sum(f1_values) / len(f1_values),
                min(f1_values),
                max(f1_values),
                max(f1_values) - min(f1_values),
                (
                    "yes"
                    if len(
                        {row["parsed_prediction"] for row in model_rows}
                    )
                    == 1
                    else "no"
                ),
            ]
        )
        rows.append(values)
    return headers, rows


INSTANCE_HEADERS = [
    "Population Index",
    "Sample ID",
    "Example ID",
    "Example Sub-ID",
    "Turn",
    "Difficulty",
    "Conflict Group",
    "Condition",
    "Reusable Mix600",
    "Model",
    "Provider",
    "Request Model",
    "Status",
    "Error",
    "Parse Failed",
    "Parse Reason",
    "Pref Exact Match",
    "Overall Precision",
    "Overall Recall",
    "Overall F1",
    "Pref Precision",
    "Pref Recall",
    "Pref F1",
    "Non-pref Precision",
    "Non-pref Recall",
    "Non-pref F1",
    "Overall TP",
    "Overall FP",
    "Overall FN",
    "Pref TP",
    "Pref FP",
    "Pref FN",
    "Non-pref TP",
    "Non-pref FP",
    "Non-pref FN",
    "GT Slot Count",
    "Pred Slot Count",
    "Ground Truth",
    "Parsed Calls",
    "Parsed Prediction",
    "Raw Output",
    "Reasoning",
    "Test Utterance",
    "Input Tokens",
    "Output Tokens",
    "Reasoning Tokens",
    "Cached Input Tokens",
    "Cache Creation Input Tokens",
    "Total Tokens",
    "Elapsed Seconds",
    "Finish Reason",
    "Reused Inference",
    "Checkpoint Attempt Count",
    "Selected Checkpoint Line",
]


def instance_row(result: Mapping[str, Any]) -> list[Any]:
    return [
        result["population_index"],
        result["sample_id"],
        result["example_id"],
        result["example_id_sub"],
        result["turn"],
        result["difficulty"],
        "conflict" if result["conflict"] else "non_conflict",
        result["condition"],
        "yes" if result["reusable_mix600"] else "no",
        result["model"],
        result["provider"],
        result["request_model"],
        result["status"],
        result["error"],
        "yes" if result["parse_failed"] else "no",
        result["parse_reason"],
        result["pref_exact_match"],
        result["overall_precision"],
        result["overall_recall"],
        result["overall_f1"],
        result["pref_precision"],
        result["pref_recall"],
        result["pref_f1"],
        result["nonpref_precision"],
        result["nonpref_recall"],
        result["nonpref_f1"],
        result["overall_tp"],
        result["overall_fp"],
        result["overall_fn"],
        result["pref_tp"],
        result["pref_fp"],
        result["pref_fn"],
        result["nonpref_tp"],
        result["nonpref_fp"],
        result["nonpref_fn"],
        result["gt_slot_count"],
        result["pred_slot_count"],
        result["ground_truth"],
        result["parsed_calls"],
        result["parsed_prediction"],
        result["raw_output"],
        result["reasoning"],
        result["test_utterance"],
        result["input_tokens"],
        result["output_tokens"],
        result["reasoning_tokens"],
        result["cached_input_tokens"],
        result["cache_creation_input_tokens"],
        result["total_tokens"],
        result["elapsed_seconds"],
        result["finish_reason"],
        "yes" if result["reused_inference"] else "no",
        result["checkpoint_attempt_count"],
        result["selected_checkpoint_line"],
    ]


def create_workbook(
    *,
    experiment_root: Path,
    output_path: Path,
    runs: Sequence[ModelRun],
    population: Sequence[dict[str, Any]],
    instance_results: Mapping[str, Sequence[dict[str, Any]]] | None,
) -> Workbook:
    wb = Workbook()
    wb.properties.title = "Experiment8 MPT_v2_0725 Vanilla LLM Results"
    wb.properties.subject = (
        "Detailed vanilla LLM evaluation with condition/conflict splits"
    )
    wb.properties.creator = "Experiment8 export_vanilla_performance_excel.py"
    wb.properties.keywords = (
        "Experiment8, MPT_v2_0725, vanilla_llm, evaluation, conflict"
    )
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    readme_sheet(
        wb,
        experiment_root=experiment_root,
        output_path=output_path,
        runs=runs,
        population=population,
    )
    dashboard_sheet(wb, runs)

    headers, rows = overall_rows(runs)
    overall_ws = write_standard_sheet(
        wb,
        name="Overall Summary",
        title_text="Overall Vanilla LLM Results",
        subtitle="Identical population; current parser; ranks are dense ranks.",
        headers=headers,
        rows=rows,
        table_name="OverallSummary",
        freeze="F5",
    )
    for row in range(5, 5 + len(runs)):
        if overall_ws.cell(row, headers.index("Parsing Failures") + 1).value == 0:
            overall_ws.cell(
                row, headers.index("Parsing Failures") + 1
            ).fill = GOOD_FILL

    headers, rows = condition_rows(runs)
    write_standard_sheet(
        wb,
        name="Condition Summary",
        title_text="Performance by Turn and Difficulty",
        subtitle="Six conditions per model: single/multi × easy/medium/hard.",
        headers=headers,
        rows=rows,
        table_name="ConditionSummary",
        freeze="E5",
    )

    headers, rows = split_rows(runs)
    split_ws = write_standard_sheet(
        wb,
        name="Split Detail",
        title_text="Performance by Turn, Difficulty, and Conflict",
        subtitle="Easy-conflict rows are retained with N=0 and blank metrics.",
        headers=headers,
        rows=rows,
        table_name="SplitDetail",
        freeze="F5",
    )
    for row in range(5, 5 + len(rows)):
        if split_ws.cell(row, headers.index("N") + 1).value == 0:
            for column in range(1, len(headers) + 1):
                split_ws.cell(row, column).fill = PatternFill(
                    "solid", fgColor=VERY_LIGHT_GRAY
                )

    matrix_sheet(wb, runs, turn="single")
    matrix_sheet(wb, runs, turn="multi")

    headers, rows = ranking_rows(runs)
    write_standard_sheet(
        wb,
        name="Rankings",
        title_text="Per-split Model Rankings",
        subtitle="Dense rank within each nonempty turn × difficulty × conflict split.",
        headers=headers,
        rows=rows,
        table_name="SplitRankings",
        freeze="F5",
    )

    headers, rows = conflict_gap_rows(runs)
    write_standard_sheet(
        wb,
        name="Conflict Gaps",
        title_text="Conflict Performance Gaps",
        subtitle="Delta = conflict score − non-conflict score for the same model and condition.",
        headers=headers,
        rows=rows,
        table_name="ConflictGaps",
        freeze="F5",
    )

    headers, rows = token_rows(runs)
    write_standard_sheet(
        wb,
        name="Token Usage",
        title_text="Token Usage and Preserved Cost",
        subtitle="Cost is shown only when exact run_summary.json cost was preserved.",
        headers=headers,
        rows=rows,
        table_name="TokenUsage",
        freeze="D5",
        color_scales=False,
    )

    headers, rows = parser_impact_rows(runs)
    parser_ws = write_standard_sheet(
        wb,
        name="Parser Impact",
        title_text="Parser Normalization Impact",
        subtitle="Exact before/after parsing-failure counts; current metrics shown for context.",
        headers=headers,
        rows=rows,
        table_name="ParserImpact",
        freeze="C5",
    )
    for row in range(5, 5 + len(rows)):
        after_column = headers.index("After Parsing Failures") + 1
        recovered_column = headers.index("Recovered Failures") + 1
        parser_ws.cell(row, after_column).fill = (
            GOOD_FILL if parser_ws.cell(row, after_column).value == 0 else WARN_FILL
        )
        if parser_ws.cell(row, recovered_column).value:
            parser_ws.cell(row, recovered_column).fill = GOOD_FILL

    headers, rows = population_composition_rows(population)
    write_standard_sheet(
        wb,
        name="Dataset Composition",
        title_text="Evaluation Population Composition",
        subtitle="Counts are taken from population/population.jsonl.",
        headers=headers,
        rows=rows,
        table_name="DatasetComposition",
        freeze="D5",
        color_scales=False,
    )

    headers, rows = run_inventory_rows(
        experiment_root, runs, expected_n=len(population)
    )
    write_standard_sheet(
        wb,
        name="Run Inventory",
        title_text="Vanilla Run Artifact Inventory",
        subtitle="Included complete runs and excluded/incomplete run directories.",
        headers=headers,
        rows=rows,
        table_name="RunInventory",
        freeze="E5",
        color_scales=False,
    )

    if instance_results is not None:
        headers, rows = pairwise_rows(runs, instance_results)
        write_standard_sheet(
            wb,
            name="Pairwise Agreement",
            title_text="Pairwise Model Agreement",
            subtitle="Prediction equality uses canonical parsed slot/value maps; wins use per-instance Overall F1.",
            headers=headers,
            rows=rows,
            table_name="PairwiseAgreement",
            freeze="C5",
        )

        headers, rows = instance_agreement_rows(
            runs, population, instance_results
        )
        write_standard_sheet(
            wb,
            name="Instance Agreement",
            title_text="Cross-model Instance Agreement",
            subtitle="One row per population instance, aligned across all four models.",
            headers=headers,
            rows=rows,
            table_name="InstanceAgreement",
            freeze="J5",
            custom_widths={
                "Ground Truth": 34,
                "Example ID": 28,
                "Example Sub-ID": 30,
            },
        )

        all_instance_rows = [
            instance_row(row)
            for run in runs
            for row in instance_results[run.key]
        ]
        write_standard_sheet(
            wb,
            name="Instance Results",
            title_text="Per-model Per-instance Audit Results",
            subtitle="Every included model × every population row, with scoring counts, outputs, and token usage.",
            headers=INSTANCE_HEADERS,
            rows=all_instance_rows,
            table_name="InstanceResults",
            freeze="J5",
            custom_widths={
                "Raw Output": 48,
                "Reasoning": 40,
                "Test Utterance": 42,
                "Ground Truth": 36,
                "Parsed Calls": 40,
                "Parsed Prediction": 40,
            },
        )

        failure_rows = [
            instance_row(row)
            for run in runs
            for row in instance_results[run.key]
            if row["parse_failed"]
        ]
        failure_ws = write_standard_sheet(
            wb,
            name="Parse Failures",
            title_text="Remaining Current Parser Failures",
            subtitle="All rows for which no function call is extractable under the updated parser.",
            headers=INSTANCE_HEADERS,
            rows=failure_rows,
            table_name="ParseFailures",
            freeze="J5",
            custom_widths={
                "Raw Output": 60,
                "Reasoning": 40,
                "Test Utterance": 42,
                "Ground Truth": 36,
            },
            color_scales=False,
        )
        for row in range(5, 5 + len(failure_rows)):
            failure_ws.cell(
                row, INSTANCE_HEADERS.index("Parse Failed") + 1
            ).fill = BAD_FILL

    return wb


def validate_workbook(
    path: Path,
    *,
    expected_n: int,
    include_instances: bool,
    expected_parse_failures: int,
) -> dict[str, Any]:
    wb = load_workbook(path, read_only=False, data_only=False)
    required = {
        "README",
        "Dashboard",
        "Overall Summary",
        "Condition Summary",
        "Split Detail",
        "Single Matrix",
        "Multi Matrix",
        "Rankings",
        "Conflict Gaps",
        "Token Usage",
        "Parser Impact",
        "Dataset Composition",
        "Run Inventory",
    }
    if include_instances:
        required.update(
            {
                "Pairwise Agreement",
                "Instance Agreement",
                "Instance Results",
                "Parse Failures",
            }
        )
    missing = required - set(wb.sheetnames)
    if missing:
        raise ValueError(f"Workbook is missing sheets: {sorted(missing)}")
    checks = {
        "overall_rows": wb["Overall Summary"].max_row - 4,
        "condition_rows": wb["Condition Summary"].max_row - 4,
        "split_rows": wb["Split Detail"].max_row - 4,
        "sheet_count": len(wb.sheetnames),
        "chart_count": sum(len(ws._charts) for ws in wb.worksheets),
        "table_count": sum(len(ws.tables) for ws in wb.worksheets),
    }
    if checks["overall_rows"] != len(MODEL_SPECS):
        raise ValueError(f"Unexpected overall row count: {checks}")
    if checks["condition_rows"] != len(MODEL_SPECS) * len(CONDITION_ORDER):
        raise ValueError(f"Unexpected condition row count: {checks}")
    if checks["split_rows"] != (
        len(MODEL_SPECS) * len(CONDITION_ORDER) * len(CONFLICT_ORDER)
    ):
        raise ValueError(f"Unexpected split row count: {checks}")
    if include_instances:
        checks["instance_agreement_rows"] = (
            wb["Instance Agreement"].max_row - 4
        )
        checks["instance_result_rows"] = wb["Instance Results"].max_row - 4
        checks["parse_failure_rows"] = wb["Parse Failures"].max_row - 4
        if checks["instance_agreement_rows"] != expected_n:
            raise ValueError(f"Unexpected instance agreement rows: {checks}")
        if checks["instance_result_rows"] != expected_n * len(MODEL_SPECS):
            raise ValueError(f"Unexpected instance result rows: {checks}")
        if checks["parse_failure_rows"] != expected_parse_failures:
            raise ValueError(f"Unexpected parse failure rows: {checks}")
    wb.close()
    return checks


def main() -> None:
    args = parse_args()
    experiment_root = args.experiment_root.resolve()
    output_path = args.output.resolve()
    population = load_population(experiment_root)
    runs = load_runs(experiment_root, expected_n=len(population))
    instance_results = (
        None
        if args.skip_instance_sheets
        else build_instance_results(runs, population)
    )
    workbook = create_workbook(
        experiment_root=experiment_root,
        output_path=output_path,
        runs=runs,
        population=population,
        instance_results=instance_results,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    checks = validate_workbook(
        output_path,
        expected_n=len(population),
        include_instances=instance_results is not None,
        expected_parse_failures=sum(
            int(run.report["overall"]["parsing_failures"]) for run in runs
        ),
    )
    result = {
        "output": str(output_path),
        "bytes": output_path.stat().st_size,
        "sha256": sha256_file(output_path),
        "population": len(population),
        "models": [run.name for run in runs],
        **checks,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
