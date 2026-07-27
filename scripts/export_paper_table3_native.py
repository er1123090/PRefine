#!/usr/bin/env python3
"""Export experiment8 native results in the paper Table 3 layout."""

from __future__ import annotations

import argparse
import csv
import gc
import importlib.metadata
import json
import math
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean, pstdev
from typing import Any, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
SCRIPT_DIR = ROOT / "scripts"
for path in (ROOT, SCRIPT_DIR):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from export_unified_vanilla_ours_performance import (  # noqa: E402
    RUNS,
    SAMPLE_SIZE,
    RunSpec,
)
from scripts.run_vanilla_batch_sample import evaluation_report  # noqa: E402
from src.evaluation import metrics as evaluation_metrics  # noqa: E402


EXPERIMENT_ROOT = (
    ROOT / "outputs" / "mpt_v2_0725_hint_no_easy_conflict"
)
DEFAULT_OUTPUT_DIR = (
    EXPERIMENT_ROOT / f"paper_table3_sample{SAMPLE_SIZE}"
)

DIFFICULTIES = ("easy", "medium", "hard")
DIFFICULTY_LABELS = {
    "easy": "Preference Recall",
    "medium": "Preference Induction",
    "hard": "Preference Transfer",
}

METRICS = (
    "CG Recall P-EM",
    "CG Recall EA-F1",
    "CG Recall OA-F1",
    "CG Induction P-EM",
    "CG Induction EA-F1",
    "CG Induction OA-F1",
    "CG Transfer P-EM",
    "CG Transfer EA-F1",
    "CG Transfer OA-F1",
    "CG Avg OA-F1",
    "CF Recall Precision",
    "CF Recall Recall",
    "CF Recall F1",
    "CF Induction Precision",
    "CF Induction Recall",
    "CF Induction F1",
    "CF Transfer Precision",
    "CF Transfer Recall",
    "CF Transfer F1",
    "CF Avg F1",
)

MODEL_ORDER = (
    "GPT-5",
    "Claude Haiku 4.5",
    "Qwen3-8B",
    "GPT-OSS-20B",
    "Gemma4-12B",
)

HEADER_GROUPS = (
    ("Preference Recall", 3),
    ("Preference Induction", 3),
    ("Preference Transfer", 3),
    ("Avg.", 1),
)

THIN_GRAY = Side(style="thin", color="A6A6A6")
MEDIUM_DARK = Side(style="medium", color="000000")
SECTION_FILL = PatternFill("solid", fgColor="E7E6E6")
AVERAGE_FILL = PatternFill("solid", fgColor="D9E1F2")
HEADER_FILL = PatternFill("solid", fgColor="D9EAD3")
SUBHEADER_FILL = PatternFill("solid", fgColor="E2F0D9")

GAIN_COLORS = {
    0: "E2F0D9",
    1: "C6E0B4",
    2: "A9D18E",
    3: "70AD47",
}
LOSS_COLORS = {
    0: "FCE4D6",
    1: "F4B084",
    2: "E26B0A",
    3: "C00000",
}


def validate_report(spec: RunSpec, report: Mapping[str, Any]) -> None:
    required = {"overall", "conditions", "condition_conflict"}
    missing = required - set(report)
    if missing:
        raise RuntimeError(
            f"{spec.key}: evaluation report missing {sorted(missing)}"
        )
    expected_n = SAMPLE_SIZE if spec.scope == "sample" else 4695
    observed_n = int(report["overall"]["n"])
    if observed_n != expected_n:
        raise RuntimeError(
            f"{spec.key}: expected N={expected_n:,}, got N={observed_n:,}"
        )


def evaluator_provenance(
    spec: RunSpec,
    predictions_path: Path,
) -> dict[str, Any]:
    source_stat = predictions_path.stat()
    return {
        "evaluator": (
            "scripts.run_vanilla_batch_sample.evaluation_report + "
            "src.evaluation.metrics"
        ),
        "dateparser_available": evaluation_metrics._DATEPARSER_AVAILABLE,
        "dateparser_version": importlib.metadata.version("dateparser"),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "run_id": spec.key,
        "source_predictions": str(predictions_path),
        "source_size_bytes": source_stat.st_size,
        "source_mtime_ns": source_stat.st_mtime_ns,
        "source_evaluation_preserved": str(
            spec.directory / "evaluation.json"
        ),
    }


def recompute_reports(output_dir: Path) -> dict[str, dict[str, Any]]:
    if not evaluation_metrics._DATEPARSER_AVAILABLE:
        raise RuntimeError(
            "dateparser is required for harmonized evaluation. "
            "Install the project requirements before exporting Table 3."
        )

    output_dir.mkdir(parents=True, exist_ok=True)
    reports: dict[str, dict[str, Any]] = {}
    for index, spec in enumerate(RUNS, start=1):
        predictions_path = spec.directory / "predictions.json"
        if not predictions_path.exists():
            raise FileNotFoundError(
                f"{spec.key}: missing {predictions_path}"
            )
        print(
            f"[harmonized evaluator {index}/{len(RUNS)}] {spec.key}",
            flush=True,
        )
        with predictions_path.open("r", encoding="utf-8") as handle:
            predictions = json.load(handle)
        if not isinstance(predictions, list):
            raise RuntimeError(
                f"{spec.key}: predictions.json must contain a list"
            )
        report = evaluation_report(predictions)
        validate_report(spec, report)
        report["harmonized_evaluation"] = evaluator_provenance(
            spec,
            predictions_path,
        )
        report_path = output_dir / f"{spec.key}.json"
        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        reports[spec.key] = report
        del predictions
        gc.collect()
    return reports


def condition_node(
    report: Mapping[str, Any],
    turn: str,
    difficulty: str,
    scope: str,
) -> Mapping[str, Any]:
    condition = f"{turn}_{difficulty}"
    if scope == "native_all":
        return report["conditions"][condition]
    if scope == "non_conflict":
        return report["condition_conflict"][condition]["non_conflict"]
    raise ValueError(f"Unsupported scope: {scope}")


def table3_metrics(
    report: Mapping[str, Any],
    scope: str,
) -> dict[str, float]:
    result: dict[str, float] = {}
    guided_oaf1: list[float] = []
    free_f1: list[float] = []
    for difficulty, short in zip(DIFFICULTIES, ("Recall", "Induction", "Transfer")):
        guided = condition_node(report, "multi", difficulty, scope)
        result[f"CG {short} P-EM"] = (
            float(guided["pref"]["exact_match_rate"]) * 100
        )
        result[f"CG {short} EA-F1"] = (
            float(guided["nonpref"]["f1"]) * 100
        )
        result[f"CG {short} OA-F1"] = (
            float(guided["overall"]["f1"]) * 100
        )
        guided_oaf1.append(result[f"CG {short} OA-F1"])

        free = condition_node(report, "single", difficulty, scope)
        result[f"CF {short} Precision"] = (
            float(free["pref"]["precision"]) * 100
        )
        result[f"CF {short} Recall"] = (
            float(free["pref"]["recall"]) * 100
        )
        result[f"CF {short} F1"] = float(free["pref"]["f1"]) * 100
        free_f1.append(result[f"CF {short} F1"])

    result["CG Avg OA-F1"] = mean(guided_oaf1)
    result["CF Avg F1"] = mean(free_f1)
    if set(result) != set(METRICS):
        raise RuntimeError(
            f"Metric mismatch: missing={set(METRICS) - set(result)}, "
            f"extra={set(result) - set(METRICS)}"
        )
    return result


def evidence_label(spec: RunSpec) -> str:
    if spec.scope == "sample":
        return (
            f"Measured on proportional N={SAMPLE_SIZE:,} stratified sample"
        )
    return "Measured on all native N=4,695 instances"


def run_rows(
    scope: str,
    reports: Mapping[str, Mapping[str, Any]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for spec in RUNS:
        report = reports[spec.key]
        row: dict[str, Any] = {
            "Run ID": spec.key,
            "Method": spec.method,
            "Memory Source": spec.memory,
            "Inference Model": spec.inference_model,
            "Inference Setting": spec.setting,
            "Provider / Backend": spec.provider,
            "Evidence": evidence_label(spec),
            "Observed N": int(report["overall"]["n"]),
            "Scope": scope,
        }
        row.update(table3_metrics(report, scope))
        rows.append(row)
    return rows


def metric_average(rows: Sequence[Mapping[str, Any]]) -> dict[str, float]:
    return {
        metric: mean(float(row[metric]) for row in rows)
        for metric in METRICS
    }


def metric_std(rows: Sequence[Mapping[str, Any]]) -> dict[str, float]:
    return {
        metric: pstdev(float(row[metric]) for row in rows)
        for metric in METRICS
    }


def paired_sections(
    rows: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    base = {
        str(row["Inference Model"]): dict(row)
        for row in rows
        if row["Method"] == "Vanilla LLM"
    }
    ours_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        if row["Method"] == "Ours-memory":
            ours_groups[str(row["Inference Model"])].append(dict(row))

    if set(base) != set(MODEL_ORDER):
        raise RuntimeError(f"Unexpected vanilla models: {sorted(base)}")
    if set(ours_groups) != set(MODEL_ORDER):
        raise RuntimeError(
            f"Unexpected ours inference models: {sorted(ours_groups)}"
        )
    for model, group in ours_groups.items():
        if len(group) != 3:
            raise RuntimeError(
                f"{model}: expected 3 memory sources, got {len(group)}"
            )

    base_rows = [base[model] for model in MODEL_ORDER]
    ours_rows: list[dict[str, Any]] = []
    for model in MODEL_ORDER:
        group = ours_groups[model]
        observed_n = min(int(row["Observed N"]) for row in group)
        averaged: dict[str, Any] = {
            "Method": "Ours-memory",
            "Memory Source": "Average of 3 memories",
            "Inference Model": model,
            "Evidence": (
                f"Measured on proportional N={observed_n:,} stratified sample"
                if observed_n != 4695
                else "Measured native N=4,695"
            ),
            "Observed N": observed_n,
            "Memory Count": len(group),
        }
        averaged.update(metric_average(group))
        ours_rows.append(averaged)

    base_average = metric_average(base_rows)
    ours_average = metric_average(ours_rows)
    gain = {
        metric: ours_average[metric] - base_average[metric]
        for metric in METRICS
    }
    return {
        "base_rows": base_rows,
        "base_average": base_average,
        "ours_rows": ours_rows,
        "ours_average": ours_average,
        "gain": gain,
        "per_memory_rows": [
            row
            for model in MODEL_ORDER
            for row in sorted(
                ours_groups[model],
                key=lambda item: str(item["Memory Source"]),
            )
        ],
        "ours_std_rows": [
            {
                "Inference Model": model,
                **metric_std(ours_groups[model]),
            }
            for model in MODEL_ORDER
        ],
    }


def delta_fill(delta: float) -> PatternFill:
    magnitude = abs(delta)
    level = 0 if magnitude < 5 else 1 if magnitude < 10 else 2 if magnitude < 20 else 3
    color = GAIN_COLORS[level] if delta >= 0 else LOSS_COLORS[level]
    return PatternFill("solid", fgColor=color)


def apply_table3_headers(sheet: Any, title: str, note: str) -> None:
    last_column = 21
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet.cell(1, 1, title)
    sheet.cell(1, 1).font = Font(size=15, bold=True, color="FFFFFF")
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor="1F4E78")
    sheet.cell(1, 1).alignment = Alignment(horizontal="center")
    sheet.row_dimensions[1].height = 25

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet.cell(2, 1, note)
    sheet.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.cell(2, 1).fill = PatternFill("solid", fgColor="D9EAF7")
    sheet.row_dimensions[2].height = 43

    sheet.merge_cells("A4:A6")
    sheet["A4"] = "Base LLM"
    sheet.merge_cells("B4:K4")
    sheet["B4"] = "Context-Guided Query (multi-turn)"
    sheet.merge_cells("L4:U4")
    sheet["L4"] = "Context-Free Query (single-turn)"

    start = 2
    for label, width in HEADER_GROUPS:
        end = start + width - 1
        sheet.merge_cells(
            start_row=5,
            start_column=start,
            end_row=5,
            end_column=end,
        )
        sheet.cell(5, start, label)
        start = end + 1
    start = 12
    for label, width in HEADER_GROUPS:
        end = start + width - 1
        sheet.merge_cells(
            start_row=5,
            start_column=start,
            end_row=5,
            end_column=end,
        )
        sheet.cell(5, start, label)
        start = end + 1

    guided_labels = (
        "P-EM",
        "EA-F1",
        "OA-F1",
        "P-EM",
        "EA-F1",
        "OA-F1",
        "P-EM",
        "EA-F1",
        "OA-F1",
        "OA-F1",
    )
    free_labels = (
        "Prec.",
        "Rec.",
        "F1",
        "Prec.",
        "Rec.",
        "F1",
        "Prec.",
        "Rec.",
        "F1",
        "F1",
    )
    for column, label in enumerate(guided_labels + free_labels, start=2):
        sheet.cell(6, column, label)

    for row in range(4, 7):
        for column in range(1, last_column + 1):
            cell = sheet.cell(row, column)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.fill = HEADER_FILL if row == 4 else SUBHEADER_FILL
            cell.border = Border(
                top=THIN_GRAY,
                bottom=THIN_GRAY,
                left=THIN_GRAY,
                right=THIN_GRAY,
            )
    sheet.row_dimensions[4].height = 25
    sheet.row_dimensions[5].height = 30
    sheet.row_dimensions[6].height = 25


def write_section_row(sheet: Any, row: int, label: str) -> None:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=21)
    cell = sheet.cell(row, 1, label)
    cell.font = Font(bold=True)
    cell.fill = SECTION_FILL
    cell.alignment = Alignment(horizontal="center")
    cell.border = Border(top=MEDIUM_DARK, bottom=THIN_GRAY)


def write_metric_row(
    sheet: Any,
    row_index: int,
    label: str,
    values: Mapping[str, float],
    *,
    average: bool = False,
    gain: bool = False,
) -> None:
    sheet.cell(row_index, 1, label)
    sheet.cell(row_index, 1).font = Font(bold=average or gain)
    sheet.cell(row_index, 1).alignment = Alignment(horizontal="left")
    for column, metric in enumerate(METRICS, start=2):
        value = float(values[metric])
        cell = sheet.cell(row_index, column, value)
        cell.number_format = "0.00"
        cell.alignment = Alignment(horizontal="right")
        if gain:
            cell.fill = delta_fill(value)
            if abs(value) >= 20 and value < 0:
                cell.font = Font(bold=True, color="FFFFFF")
        elif average:
            cell.fill = AVERAGE_FILL
            cell.font = Font(bold=True)
        cell.border = Border(bottom=THIN_GRAY)
    if average:
        sheet.cell(row_index, 1).fill = AVERAGE_FILL
    if gain:
        sheet.cell(row_index, 1).fill = SECTION_FILL


def write_table3_sheet(
    workbook: Workbook,
    name: str,
    sections: Mapping[str, Any],
    note: str,
) -> None:
    sheet = workbook.create_sheet(name)
    apply_table3_headers(
        sheet,
        "Experiment8 Results — Paper Table 3 Layout",
        note,
    )
    write_section_row(sheet, 7, "BASE PROMPTING: Full-dialogue context")
    base_row_indices: dict[str, int] = {}
    for index, record in enumerate(sections["base_rows"], start=8):
        model = str(record["Inference Model"])
        base_row_indices[model] = index
        write_metric_row(sheet, index, model, record)
    write_metric_row(
        sheet,
        13,
        "Average",
        sections["base_average"],
        average=True,
    )

    write_section_row(
        sheet,
        14,
        "OURS MEMORY: Average over Qwen3-8B, Gemma4-12B, and GPT-OSS-20B memories",
    )
    ours_row_indices: dict[str, int] = {}
    for index, record in enumerate(sections["ours_rows"], start=15):
        model = str(record["Inference Model"])
        label = model
        if int(record["Observed N"]) != 4695:
            label += "†"
        ours_row_indices[model] = index
        write_metric_row(sheet, index, label, record)
        if int(record["Observed N"]) != 4695:
            sheet.cell(index, 1).comment = Comment(
                "Measured on the completed proportional "
                f"N={int(record['Observed N']):,} stratified Batch sample.",
                "Codex",
            )
        for column, metric in enumerate(METRICS, start=2):
            base_value = float(
                sections["base_rows"][MODEL_ORDER.index(model)][metric]
            )
            delta = float(record[metric]) - base_value
            sheet.cell(index, column).fill = delta_fill(delta)

    write_metric_row(
        sheet,
        20,
        "Avg. Gain (%p)",
        sections["gain"],
        gain=True,
    )
    write_metric_row(
        sheet,
        21,
        "Average",
        sections["ours_average"],
        average=True,
    )

    model_rows = list(base_row_indices.values()) + list(
        ours_row_indices.values()
    )
    for column in range(2, 22):
        ranked = sorted(
            (
                (float(sheet.cell(row, column).value), row)
                for row in model_rows
            ),
            reverse=True,
        )
        best_value = ranked[0][0]
        second_value = next(
            (value for value, _ in ranked if value < best_value),
            best_value,
        )
        for value, row in ranked:
            cell = sheet.cell(row, column)
            if math.isclose(value, best_value):
                cell.font = Font(bold=True)
            elif math.isclose(value, second_value):
                cell.font = Font(underline="single")

    sheet.freeze_panes = "B7"
    sheet.column_dimensions["A"].width = 24
    for column in range(2, 22):
        sheet.column_dimensions[get_column_letter(column)].width = 10
    sheet.auto_filter.ref = "A6:U21"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_title_rows = "4:6"
    sheet.print_area = "A1:U21"


def flat_table_rows(sections: Mapping[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for record in sections["base_rows"]:
        rows.append(
            {
                "Section": "Base prompting",
                "Base LLM": record["Inference Model"],
                "Evidence": record["Evidence"],
                **{metric: record[metric] for metric in METRICS},
            }
        )
    rows.append(
        {
            "Section": "Base prompting",
            "Base LLM": "Average",
            "Evidence": "Macro-average across five inference models",
            **sections["base_average"],
        }
    )
    for record in sections["ours_rows"]:
        rows.append(
            {
                "Section": "Ours-memory",
                "Base LLM": record["Inference Model"],
                "Evidence": record["Evidence"],
                **{metric: record[metric] for metric in METRICS},
            }
        )
    rows.append(
        {
            "Section": "Ours-memory",
            "Base LLM": "Avg. Gain (%p)",
            "Evidence": "Ours average minus vanilla average",
            **sections["gain"],
        }
    )
    rows.append(
        {
            "Section": "Ours-memory",
            "Base LLM": "Average",
            "Evidence": "Macro-average across five inference models",
            **sections["ours_average"],
        }
    )
    return rows


def write_flat_sheet(
    workbook: Workbook,
    name: str,
    title: str,
    rows: Sequence[Mapping[str, Any]],
) -> None:
    sheet = workbook.create_sheet(name)
    headers = list(rows[0])
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(headers),
    )
    sheet.cell(1, 1, title)
    sheet.cell(1, 1).font = Font(size=14, bold=True, color="FFFFFF")
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor="1F4E78")
    header_row = 3
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    for row_index, row in enumerate(rows, start=4):
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row_index, column, row[header])
            if header in METRICS:
                cell.number_format = "0.00"
    last_row = 3 + len(rows)
    table = Table(
        displayName=f"{name.replace('_', '')}Table",
        ref=f"A3:{get_column_letter(len(headers))}{last_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
        showFirstColumn=False,
        showLastColumn=False,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = "D4"
    for column, header in enumerate(headers, start=1):
        width = 38 if header == "Evidence" else 20 if column <= 3 else 15
        sheet.column_dimensions[get_column_letter(column)].width = width


def per_memory_flat_rows(
    sections: Mapping[str, Any],
) -> list[dict[str, Any]]:
    return [
        {
            "Memory Source": row["Memory Source"],
            "Inference Model": row["Inference Model"],
            "Inference Setting": row["Inference Setting"],
            "Evidence": row["Evidence"],
            "Observed N": row["Observed N"],
            **{metric: row[metric] for metric in METRICS},
        }
        for row in sections["per_memory_rows"]
    ]


def raw_long_rows(
    scope_rows: Mapping[str, Sequence[Mapping[str, Any]]],
    reports: Mapping[str, Mapping[str, Any]],
) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for scope, rows in scope_rows.items():
        for row in rows:
            report = reports[str(row["Run ID"])]
            for turn in ("multi", "single"):
                for difficulty in DIFFICULTIES:
                    node = condition_node(report, turn, difficulty, scope)
                    output.append(
                        {
                            "Run ID": row["Run ID"],
                            "Method": row["Method"],
                            "Memory Source": row["Memory Source"],
                            "Inference Model": row["Inference Model"],
                            "Scope": scope,
                            "Turn": turn,
                            "Paper Query Setting": (
                                "Context-Guided"
                                if turn == "multi"
                                else "Context-Free"
                            ),
                            "Difficulty": difficulty,
                            "Paper Preference Type": DIFFICULTY_LABELS[
                                difficulty
                            ],
                            "N": int(node["n"]),
                            "P-EM": (
                                float(node["pref"]["exact_match_rate"]) * 100
                            ),
                            "Preference Precision": (
                                float(node["pref"]["precision"]) * 100
                            ),
                            "Preference Recall": (
                                float(node["pref"]["recall"]) * 100
                            ),
                            "Preference F1": (
                                float(node["pref"]["f1"]) * 100
                            ),
                            "EA-F1 / Non-pref F1": (
                                float(node["nonpref"]["f1"]) * 100
                            ),
                            "OA-F1 / Overall F1": (
                                float(node["overall"]["f1"]) * 100
                            ),
                            "API Errors": int(node["api_errors"]),
                            "Parsing Failures": int(
                                node["parsing_failures"]
                            ),
                        }
                    )
    return output


def write_mapping_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Metric_Mapping")
    rows = (
        ("Paper axis", "Experiment8 source", "Interpretation"),
        (
            "Evaluator",
            "scripts.run_vanilla_batch_sample.evaluation_report + "
            "src.evaluation.metrics",
            "Every vanilla and ours-memory predictions.json is recomputed "
            "through this exact evaluator with dateparser enabled.",
        ),
        (
            "Date/time normalization",
            "dateparser",
            "Date/time slot values use one shared normalizer for all 20 runs.",
        ),
        (
            "Context-Guided Query",
            "multi-turn, query=hint",
            "Current-session dialogue provides explicit argument context.",
        ),
        (
            "Context-Free Query",
            "single-turn, query=hint",
            "No multi-turn current-session dialogue; preference completion is isolated.",
        ),
        (
            "Preference Recall",
            "pref_type=easy",
            "Direct reuse of recurring in-domain preference evidence.",
        ),
        (
            "Preference Induction",
            "pref_type=medium",
            "Cross-session evidence aggregation within observed domains.",
        ),
        (
            "Preference Transfer",
            "pref_type=hard",
            "Preference application to a domain absent from evidence.",
        ),
        (
            "P-EM",
            "pref.exact_match_rate",
            "Exact match on preference-driven arguments.",
        ),
        (
            "EA-F1",
            "nonpref.f1",
            "F1 on explicitly specified, non-preference arguments.",
        ),
        (
            "OA-F1",
            "overall.f1",
            "F1 over all API arguments.",
        ),
        (
            "Context-Free Prec./Rec./F1",
            "pref.precision / pref.recall / pref.f1",
            "Preference-driven argument completion metrics.",
        ),
        (
            "Ours-memory row",
            "Arithmetic mean across 3 memory sources",
            "Qwen3-8B, Gemma4-12B, and GPT-OSS-20B construction memories.",
        ),
        (
            "Native_All scope",
            "conditions",
            "All available native-hint rows for each run, including medium/hard "
            f"conflict cases. API ours-memory runs use the stratified N={SAMPLE_SIZE:,} subset.",
        ),
        (
            "NonConflict scope",
            "condition_conflict.*.non_conflict",
            "Conflict cases removed for closer comparison to the paper's original table.",
        ),
        (
            "†",
            "Ours GPT-5 / Claude Haiku 4.5",
            f"Measured rates from the completed proportional N={SAMPLE_SIZE:,} "
            "stratified Batch sample.",
        ),
    )
    for row_index, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_index == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="4472C4")
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 45
    sheet.column_dimensions["C"].width = 75
    sheet.freeze_panes = "A2"


def write_readme_sheet(workbook: Workbook, output_path: Path) -> None:
    sheet = workbook.create_sheet("README")
    rows = (
        ("Workbook", output_path.name),
        ("Generated", datetime.now(timezone.utc).isoformat()),
        ("Paper", str(ROOT / "_paper" / "8. Latent_Preference_Modeling_for_Cross_Session_Personalized_Tool_Calling.pdf")),
        ("Dataset", str(ROOT / "data" / "MPT_v2_0725.json")),
        (
            "Primary sheet",
            "Table3_Native_All is the paper-style table. Full runs use N=4,695; "
            f"GPT-5/Claude ours-memory runs use the completed stratified N={SAMPLE_SIZE:,} sample.",
        ),
        (
            "Paper-comparable sheet",
            "Table3_NonConflict removes medium/hard conflict rows.",
        ),
        (
            "Ours aggregation",
            "Each inference-model row is the arithmetic mean across three memory-construction models.",
        ),
        (
            "Evidence boundary",
            "Nine ours-memory runs and all five vanilla runs are measured on N=4,695. "
            f"The six GPT-5/Claude ours-memory runs are measured on the proportional "
            f"N={SAMPLE_SIZE:,} stratified sample.",
        ),
        (
            "Evaluator consistency",
            "All 20 predictions.json files were recomputed in this export with "
            "the same evaluator and date/time normalization environment. Original "
            "per-run evaluation.json files were not modified.",
        ),
        (
            "Easy conflict",
            "Easy-conflict instances are excluded from the experiment population.",
        ),
        (
            "Percent scale",
            "All paper metrics are stored on a 0–100 scale and displayed to two decimals.",
        ),
    )
    for row_index, (key, value) in enumerate(rows, start=1):
        sheet.cell(row_index, 1, key).font = Font(bold=True)
        sheet.cell(row_index, 1).fill = PatternFill(
            "solid", fgColor="D9EAF7"
        )
        sheet.cell(row_index, 2, value).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        sheet.row_dimensions[row_index].height = 34
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 115


def write_csv(path: Path, rows: Sequence[Mapping[str, Any]]) -> None:
    if not rows:
        raise ValueError("CSV rows must not be empty")
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def latex_escape(value: str) -> str:
    return (
        value.replace("\\", r"\textbackslash{}")
        .replace("&", r"\&")
        .replace("%", r"\%")
        .replace("_", r"\_")
        .replace("#", r"\#")
    )


def latex_row(
    label: str,
    values: Mapping[str, float],
    *,
    label_is_latex: bool = False,
) -> str:
    metrics = " & ".join(f"{float(values[key]):.2f}" for key in METRICS)
    rendered_label = label if label_is_latex else latex_escape(label)
    return f"{rendered_label} & {metrics} \\\\"


def write_latex(path: Path, sections: Mapping[str, Any]) -> None:
    lines = [
        r"% Requires: \usepackage{booktabs,multirow,graphicx}",
        r"\begin{table*}[t]",
        r"\centering",
        r"\scriptsize",
        r"\resizebox{\textwidth}{!}{%",
        r"\begin{tabular}{l*{20}{r}}",
        r"\toprule",
        (
            r"& \multicolumn{10}{c}{Context-Guided Query} "
            r"& \multicolumn{10}{c}{Context-Free Query} \\"
        ),
        (
            r"Base LLM & \multicolumn{3}{c}{Pref. Recall} "
            r"& \multicolumn{3}{c}{Pref. Induction} "
            r"& \multicolumn{3}{c}{Pref. Transfer} & Avg. "
            r"& \multicolumn{3}{c}{Pref. Recall} "
            r"& \multicolumn{3}{c}{Pref. Induction} "
            r"& \multicolumn{3}{c}{Pref. Transfer} & Avg. \\"
        ),
        (
            r"& P-EM & EA-F1 & OA-F1 & P-EM & EA-F1 & OA-F1 "
            r"& P-EM & EA-F1 & OA-F1 & OA-F1 "
            r"& Prec. & Rec. & F1 & Prec. & Rec. & F1 "
            r"& Prec. & Rec. & F1 & F1 \\"
        ),
        r"\midrule",
        r"\multicolumn{21}{c}{\textsc{Base Prompting}: Full-dialogue context} \\",
    ]
    for row in sections["base_rows"]:
        lines.append(latex_row(str(row["Inference Model"]), row))
    lines.append(latex_row("Average", sections["base_average"]))
    lines.extend(
        [
            r"\midrule",
            (
                r"\multicolumn{21}{c}{\textsc{Ours Memory}: "
                r"Average over three construction models} \\"
            ),
        ]
    )
    for row in sections["ours_rows"]:
        label = latex_escape(str(row["Inference Model"]))
        if int(row["Observed N"]) != 4695:
            label += r"$^\dagger$"
        lines.append(latex_row(label, row, label_is_latex=True))
    lines.append(latex_row("Avg. Gain (%p)", sections["gain"]))
    lines.append(latex_row("Average", sections["ours_average"]))
    lines.extend(
        [
            r"\bottomrule",
            r"\end{tabular}%",
            r"}",
            (
                r"\caption{Experiment8 results in the layout of Table 3. "
                r"$^\dagger$ denotes measured rates from the proportional "
                f"$N={SAMPLE_SIZE:,}$ stratified sample."
                r"}"
            ),
            r"\label{tab:experiment8-native-table3}",
            r"\end{table*}",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def build_outputs(output_dir: Path) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    harmonized_dir = output_dir / "harmonized_evaluations"
    reports = recompute_reports(harmonized_dir)
    native_rows = run_rows("native_all", reports)
    nonconflict_rows = run_rows("non_conflict", reports)
    native_sections = paired_sections(native_rows)
    nonconflict_sections = paired_sections(nonconflict_rows)

    workbook_path = (
        output_dir / f"experiment8_table3_sample{SAMPLE_SIZE}.xlsx"
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    write_table3_sheet(
        workbook,
        "Table3_Native_All",
        native_sections,
        (
            "Same metric layout as paper Table 3. Native hint population includes "
            "medium/hard conflict cases; easy-conflict is excluded. Ours rows are "
            "averaged across 3 memory-construction models. "
            f"† marks measured N={SAMPLE_SIZE:,} stratified Batch results."
        ),
    )
    write_table3_sheet(
        workbook,
        "Table3_NonConflict",
        nonconflict_sections,
        (
            "Paper-comparable non-conflict view. Multi-turn maps to Context-Guided, "
            "single-turn to Context-Free, and easy/medium/hard to Recall/Induction/"
            "Transfer. Ours rows average 3 memory sources; "
            f"† marks measured N={SAMPLE_SIZE:,} stratified Batch results."
        ),
    )
    write_flat_sheet(
        workbook,
        "Native_Flat",
        "Flat Table 3 Values — Native All",
        flat_table_rows(native_sections),
    )
    write_flat_sheet(
        workbook,
        "NonConflict_Flat",
        "Flat Table 3 Values — Non-conflict",
        flat_table_rows(nonconflict_sections),
    )
    write_flat_sheet(
        workbook,
        "Per_Memory_Native",
        "Ours-memory 3×5 Native Results Before Memory-source Averaging",
        per_memory_flat_rows(native_sections),
    )
    write_flat_sheet(
        workbook,
        "Raw_Long",
        "Run × Query Setting × Preference Type — Traceable Raw Metrics",
        raw_long_rows(
            {
                "native_all": native_rows,
                "non_conflict": nonconflict_rows,
            },
            reports,
        ),
    )
    write_mapping_sheet(workbook)
    write_readme_sheet(workbook, workbook_path)
    workbook.active = 0
    workbook.save(workbook_path)

    native_flat = flat_table_rows(native_sections)
    per_memory_flat = per_memory_flat_rows(native_sections)
    native_csv = output_dir / f"experiment8_table3_sample{SAMPLE_SIZE}.csv"
    per_memory_csv = (
        output_dir
        / f"experiment8_table3_sample{SAMPLE_SIZE}_per_memory.csv"
    )
    latex_path = output_dir / f"experiment8_table3_sample{SAMPLE_SIZE}.tex"
    write_csv(native_csv, native_flat)
    write_csv(per_memory_csv, per_memory_flat)
    write_latex(latex_path, native_sections)

    verified = load_workbook(workbook_path, read_only=True, data_only=True)
    expected_sheets = (
        "Table3_Native_All",
        "Table3_NonConflict",
        "Native_Flat",
        "NonConflict_Flat",
        "Per_Memory_Native",
        "Raw_Long",
        "Metric_Mapping",
        "README",
    )
    if tuple(verified.sheetnames) != expected_sheets:
        raise RuntimeError(
            f"Unexpected workbook sheets: {verified.sheetnames}"
        )
    if verified["Table3_Native_All"].max_row != 21:
        raise RuntimeError("Table3_Native_All must end at row 21")
    if verified["Per_Memory_Native"].max_row != 18:
        raise RuntimeError("Per_Memory_Native must contain 15 rows")
    if verified["Raw_Long"].max_row != 243:
        raise RuntimeError("Raw_Long must contain 240 rows")
    verified.close()

    return {
        "workbook": str(workbook_path),
        "native_csv": str(native_csv),
        "per_memory_csv": str(per_memory_csv),
        "latex": str(latex_path),
        "harmonized_evaluations": str(harmonized_dir),
        "harmonized_run_count": len(reports),
        "dateparser_version": importlib.metadata.version("dateparser"),
        "vanilla_model_rows": len(native_sections["base_rows"]),
        "ours_averaged_model_rows": len(native_sections["ours_rows"]),
        "ours_per_memory_rows": len(native_sections["per_memory_rows"]),
        "full_native_ours_runs": sum(
            row["Observed N"] == 4695
            for row in native_sections["per_memory_rows"]
        ),
        "stratified_sample_ours_runs": sum(
            row["Observed N"] != 4695
            for row in native_sections["per_memory_rows"]
        ),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
    )
    args = parser.parse_args()
    summary = build_outputs(args.output_dir.resolve())
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
